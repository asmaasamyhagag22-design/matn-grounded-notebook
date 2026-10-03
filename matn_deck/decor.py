"""Native-shape visuals the renderer draws for filled slots only (Step 2c): cards, accent strips, progress bars,
pictogram icons, timeline line and dots, numbered markers' cards and the section tracker.

Every fill is a theme color reference (accent1 with the template's own accent shade, or tx1 with transparency), so
each template keeps its palette and dark/light mode. Shapes carry no theme style, so no shadows or outlines appear.
"""
from __future__ import annotations

from dataclasses import dataclass

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE

from .scoring import EMU_PER_PT, Rect

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def tint(slot: str, alpha: float) -> str:
    return f'<a:schemeClr xmlns:a="{_A}" val="{slot}"><a:alpha val="{round(alpha * 100000)}"/></a:schemeClr>'


CARD_ALPHA = 0.07  # tx1 over the background: reads on light and dark themes alike
HERO_ALPHA = 0.12
MUTED_ALPHA = 0.2
# kind -> (shape, fill: "accent" or (slot, alpha), behind the text?)
_KINDS = {
    "card": (MSO_SHAPE.ROUNDED_RECTANGLE, ("tx1", CARD_ALPHA), True),
    "hero_card": (MSO_SHAPE.ROUNDED_RECTANGLE, ("accent1", HERO_ALPHA), True),
    "strip": (MSO_SHAPE.RECTANGLE, "accent", True),
    "bar_track": (MSO_SHAPE.ROUNDED_RECTANGLE, ("tx1", MUTED_ALPHA), False),
    "bar_fill": (MSO_SHAPE.ROUNDED_RECTANGLE, "accent", False),
    "icon_on": (MSO_SHAPE.OVAL, "accent", False),
    "icon_off": (MSO_SHAPE.OVAL, ("tx1", MUTED_ALPHA), False),
    "timeline_line": (MSO_SHAPE.RECTANGLE, ("tx1", 0.3), True),
    "timeline_dot": (MSO_SHAPE.OVAL, "accent", False),
    "tracker_on": (MSO_SHAPE.OVAL, "accent", False),
    "tracker_off": (MSO_SHAPE.OVAL, ("tx1", 0.3), False),
}
CARD_RADIUS_PT = 10


@dataclass(frozen=True)
class Decor:
    kind: str
    rect: Rect
    name: str


def color_xml(kind: str, accent_xml: str) -> str:
    fill = _KINDS[kind][1]
    return accent_xml if fill == "accent" else tint(*fill)


def draw(slide, items: list[Decor], accent_xml: str) -> list:
    """Add each decoration as a native autoshape; cards and other backdrops go behind every placeholder."""
    tree = slide.shapes._spTree
    shapes = []
    behind = 2  # after nvGrpSpPr and grpSpPr
    for item in items:
        kind, fill, back = _KINDS[item.kind]
        r = item.rect
        shape = slide.shapes.add_shape(kind, int(r.x), int(r.y), max(1, int(r.w)), max(1, int(r.h)))
        shape.name = item.name
        element = shape._element
        style = element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            element.remove(style)
        sp_pr = element.spPr
        for child in list(sp_pr):
            if etree.QName(child).localname in ("solidFill", "noFill", "gradFill", "ln"):
                sp_pr.remove(child)
        solid = etree.SubElement(sp_pr, f"{{{_A}}}solidFill")
        solid.append(etree.fromstring(color_xml(item.kind, accent_xml)))
        line = etree.SubElement(sp_pr, f"{{{_A}}}ln")
        etree.SubElement(line, f"{{{_A}}}noFill")
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            radius = CARD_RADIUS_PT * EMU_PER_PT if item.kind in ("card", "hero_card") else min(r.w, r.h) / 2
            shape.adjustments[0] = min(0.5, radius / max(1.0, min(r.w, r.h)))
        if shape.has_text_frame:
            shape.text_frame.text = ""
        if back:
            tree.remove(element)
            tree.insert(behind, element)
            behind += 1
        shapes.append(shape)
    return shapes
