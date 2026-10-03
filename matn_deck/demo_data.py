"""Read-only data for the demo dashboard (app/dashboard.py): manifest, gallery, receipts, scores, stats.

No Streamlit and no PDF library here, so it is testable and the dashboard never imports PyMuPDF. Everything is read
from files that already exist (out/demo/, design/templates/, out/gallery/, reports/), so the dashboard works
offline even when PowerPoint and Ollama are unavailable.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from matn_contracts.conversion import numeric_tokens

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "demo"
TEMPLATES = ROOT / "design" / "templates"
GALLERY_JSON = TEMPLATES / "src" / "gallery.json"
SCORES_MD = ROOT / "out" / "gallery" / "scores.md"
TEST_RUN = OUT / "last_test_run.json"
HANDOFF = ROOT / "reports" / "HANDOFF.md"
_FACT_ID = re.compile(r"fact:[0-9a-f]{16}")
_PART = re.compile(r"\s\(\d+/\d+\)$")


def rel(path: str | Path) -> Path:
    return ROOT / path


def load_manifest() -> dict:
    """out/demo/manifest.json, rebuilt from disk when it is missing."""
    path = OUT / "manifest.json"
    if path.is_file():
        return json.loads(path.read_text(encoding="utf-8"))
    from .demo_build import manifest  # noqa: PLC0415 - only when the cache has no manifest yet

    return manifest()


def load_facts(doc_key: str) -> dict:
    return json.loads((OUT / "facts" / f"{doc_key}.json").read_text(encoding="utf-8"))


def load_plan(doc_key: str) -> dict:
    path = OUT / "decks" / doc_key / "plan.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


# ---- gallery --------------------------------------------------------------------------------------------------------

def gallery_entries() -> list[dict]:
    """Gallery templates that passed QA (not excluded, files present): dark first, then the default, then by id."""
    data = json.loads(GALLERY_JSON.read_text(encoding="utf-8"))
    default = data["default"]
    entries = []
    for entry in data["templates"]:
        folder = TEMPLATES / entry["id"]
        if entry.get("excluded") or not (folder / "template.pptx").is_file() or not (folder / "theme.json").is_file():
            continue
        theme = json.loads((folder / "theme.json").read_text(encoding="utf-8"))
        entries.append({**entry, "mode": entry.get("mode") or theme.get("mode", "light"),
                        "name_ar": entry.get("name_ar") or theme.get("name_ar"),
                        "thumbnail": folder / "thumbnail.png", "template": folder / "template.pptx",
                        "default": entry["id"] == default})
    return sorted(entries, key=lambda e: (e["mode"] != "dark", not e["default"], e["family"], e["id"]))


def excluded_templates() -> list[dict]:
    data = json.loads(GALLERY_JSON.read_text(encoding="utf-8"))
    return [{"id": e["id"], "reason": e["excluded"]} for e in data["templates"] if e.get("excluded")]


# ---- decks ----------------------------------------------------------------------------------------------------------

@dataclass(frozen=True)
class SlideInfo:
    number: int
    title: str
    texts: tuple[str, ...]
    fact_ids: tuple[str, ...]
    layout: str


def deck_slides(pptx: Path) -> list[SlideInfo]:
    """Each slide's title, texts and cited fact IDs (from the provenance lines the renderer writes in the notes)."""
    prs = Presentation(str(pptx))
    out = []
    for number, slide in enumerate(prs.slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        title = " ".join(title.replace("", " ").split())
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append(SlideInfo(number, title, tuple(texts), tuple(dict.fromkeys(_FACT_ID.findall(notes))),
                             slide.slide_layout.name))
    return out


def deck_numbers(pptx: Path) -> list[str]:
    """Every number shown on the content slides, as written, sorted (sources slides, page chips and "(1/2)" part
    markers are left out: they depend on how a template splits slides, not on the content)."""
    numbers = []
    for slide in deck_slides(pptx):
        if slide.layout.startswith("matn:sources"):
            continue
        for text in slide.texts:
            if re.fullmatch(r"p\.\s[\d,\s]+(\(table\))?", text.strip()) or re.fullmatch(r"\d{1,2}", text.strip()):
                continue  # source chip, ordinal or marker
            text = _PART.sub("", text.strip())
            numbers += [m.group().strip() for m in numeric_tokens(text)]
    return sorted(numbers)


def deck_fact_ids(pptx: Path) -> list[str]:
    return sorted({fid for slide in deck_slides(pptx) for fid in slide.fact_ids})


def receipts_for(doc_key: str, fact_ids: list[str] | tuple[str, ...]) -> list[dict]:
    """The receipt of each cited fact: claim, number, quote, pages and the pre-rendered page images."""
    data = load_facts(doc_key)
    facts = {f["fact_id"]: f for f in data["facts"]}
    out = []
    for fid in fact_ids:
        fact, receipt = facts.get(fid), data["receipts"].get(fid)
        if fact is None or receipt is None:
            continue
        number = fact.get("number_value")
        out.append({
            "fact_id": fid, "kind": fact["kind"], "status": fact["status"], "claim": fact["statement"],
            "value": number["value_as_written"] if number else None, "unit": number["unit"] if number else None,
            "label": number["label"] if number else None, "quote": receipt["quote"],
            "page_label": receipt["page_label"], "physical_page": receipt["physical_page"],
            "section": " › ".join(receipt["section_path"]),
            "crop_png": rel(receipt["crop_png"]) if receipt.get("crop_png") else None,
            "page_png": rel(receipt["page_png"]) if receipt.get("page_png") else None,
        })
    return out


# ---- quality and stats ----------------------------------------------------------------------------------------------

def scores_table() -> list[dict]:
    """The first markdown table of out/gallery/scores.md as rows (empty when the file is missing)."""
    if not SCORES_MD.is_file():
        return []
    lines = [line for line in SCORES_MD.read_text(encoding="utf-8").splitlines()]
    start = next((i for i, line in enumerate(lines) if line.startswith("| Template")), None)
    if start is None:
        return []
    header = [c.strip() for c in lines[start].strip("|").split("|")]
    rows = []
    for line in lines[start + 2:]:
        if not line.startswith("|"):
            break
        rows.append(dict(zip(header, [c.strip() for c in line.strip("|").split("|")])))
    return rows


EVOLUTION = [  # paths under out/gallery/; gallery QA overwrites the live sheets, so 2d is read from its snapshot
    ("2b", "contact_sheet_v2.png", "Step 2b: أول gallery، القوالب من Office themes. الأغلفة حلوة والجوّه باهت"),
    ("2c", "contact_sheet_v3.png", "Step 2c: generate-and-score: dividers، KPI، charts، pictograms"),
    ("2c", "before_after.png", "Step 2c: قبل وبعد"),
    ("2c.1", "history/2d/before_after_v4.png", "Step 2c.1: قبل وبعد (fill من الحبر، cards بتحضن المحتوى، 3 مقاسات)"),
    ("2d", "history/2d/contact_sheet_creative.png", "Step 2d: العيلة الإبداعية (ليل النيل، بردي، مشربية)"),
    ("2d", "history/2d/families.png", "Step 2d: العائلات جنب بعض"),
    ("polish", "polish_before_after.png",
     "تلميع الليلة: contrast للبروجيكتور في القوالب الداكنة، و art أوضح في ليل النيل وبردي، و Modulo"),
    ("polish", "contact_sheet_creative.png", "بعد التلميع: العيلة الإبداعية"),
]


def preview(path: Path, width: int = 1600) -> Path:
    """A downscaled JPEG copy under out/demo/previews/ (the contact sheets are up to 5 MB); the original if it fails."""
    target = OUT / "previews" / f"{path.parent.name}_{path.stem}_{width}.jpg"
    try:
        if not target.is_file() or target.stat().st_mtime < path.stat().st_mtime:
            from PIL import Image  # noqa: PLC0415

            target.parent.mkdir(parents=True, exist_ok=True)
            with Image.open(path) as image:
                image = image.convert("RGB")
                if image.width > width:
                    image = image.resize((width, round(image.height * width / image.width)))
                image.save(target, "JPEG", quality=85)
        return target
    except OSError:
        return path


def evolution_images() -> list[tuple[str, Path, str]]:
    folder = ROOT / "out" / "gallery"
    return [(step, preview(folder / name), caption) for step, name, caption in EVOLUTION if (folder / name).is_file()]


def last_test_run() -> dict:
    """The last recorded full-suite result: out/demo/last_test_run.json, else the last one in HANDOFF.md."""
    if TEST_RUN.is_file():
        return json.loads(TEST_RUN.read_text(encoding="utf-8"))
    if HANDOFF.is_file():
        found = re.findall(r"`(\d+) passed(?:, (\d+) skipped)? in ([\d.]+)s`", HANDOFF.read_text(encoding="utf-8"))
        if found:
            passed, skipped, seconds = found[-1]
            return {"passed": int(passed), "skipped": int(skipped or 0), "failed": 0, "seconds": float(seconds),
                    "source": "reports/HANDOFF.md"}
    return {}


def stats() -> dict:
    manifest = load_manifest()
    docs = manifest.get("documents", {})
    return {
        "templates": len(gallery_entries()),
        "documents": len(docs),
        "verified_facts": sum(d.get("verified", 0) for d in docs.values()),
        "number_facts": sum(d.get("number_facts", 0) for d in docs.values()),
        "decks": sum(len(d.get("decks", {})) for d in docs.values()),
        "tests": last_test_run(),
    }
