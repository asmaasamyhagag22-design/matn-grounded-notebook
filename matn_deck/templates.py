"""Template contract (docs/design/template-contract.md), validator and gallery.

A template is a .pptx whose layouts are named "matn:<slide_type>:<variant>" and whose placeholders carry contract
names. `validate_template` checks a file against the contract; `gallery()` lists the templates that pass.
"""
from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
import re

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from matn_contracts.enums import SlideType

GALLERY_DIR = Path(__file__).resolve().parents[1] / "design" / "templates"
_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _items(n: int) -> tuple[tuple[str, ...], tuple[str, ...]]:
    return tuple(f"item_{i}" for i in range(1, n + 1)), tuple(f"item_{i}_title" for i in range(1, n + 1))


def _kpis(n: int) -> tuple[str, ...]:
    return tuple(f"{kind}_{i}" for i in range(1, n + 1) for kind in ("kpi_value", "kpi_label"))


_ITEMS_4, _TITLES_4 = _items(4)
_ITEMS_3, _TITLES_3 = _items(3)
_ITEMS_2, _TITLES_2 = _items(2)

def _markers(n: int) -> tuple[str, ...]:
    return tuple(f"marker_{i}" for i in range(1, n + 1))


_HERO_PANEL = ("kpi_value_1", "kpi_label_1", "quote", "attribution")

# layout name -> (required placeholders, optional placeholders), in contract order (template-contract.md §3).
LAYOUTS: dict[str, tuple[tuple[str, ...], tuple[str, ...]]] = {
    "matn:title:default": (("title",), ("subtitle",)),
    "matn:executive_summary:default": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:kpi:big-number": (("title", *_kpis(2), "source"), ("body",)),
    "matn:kpi:cards": (("title", *_kpis(4), "source"), ("body",)),
    "matn:finding:split-60-40": (("title", "body", "quote", "source"), ("attribution",)),
    "matn:finding:cards-3": (("title", *_ITEMS_3, "source"), _TITLES_3),
    "matn:finding:two-col": (("title", *_ITEMS_2, "source"), _TITLES_2),
    "matn:chart:hero": (("title", "chart", "source"), ("body",)),
    "matn:implications:default": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:limitations:default": (("title", "body", "source"), ()),
    "matn:sources:default": (("title", "body"), ()),
    "matn:section_divider:default": (("title", "section_number"), ("subtitle",)),
    "matn:closing:default": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:evidence_appendix:default": (("title", "body", "source"), ()),
    # Step 2c variants (template-contract.md §3, second table)
    "matn:executive_summary:grid-2x2": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:kpi:hero-number": (("title", *_kpis(1), "source"), ("body",)),
    "matn:kpi:bento": (("title", *_kpis(4), "source"), ("body",)),
    "matn:kpi:pictogram": (("title", *_kpis(2), "source"), ("body",)),
    "matn:kpi:progress": (("title", *_kpis(4), "source"), ("body",)),
    "matn:finding:two-split": (("title", *_ITEMS_2, *_markers(2), "source"), _TITLES_2),
    "matn:finding:grid-2x2": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:finding:timeline": (("title", *_ITEMS_4, "source"), _TITLES_4),
    "matn:finding:single-hero": (("title", "item_1", "source"), ("item_1_title", *_HERO_PANEL)),
    "matn:implications:cards-3": (("title", *_ITEMS_3, "source"), _TITLES_3),
    "matn:implications:two-split": (("title", *_ITEMS_2, *_markers(2), "source"), _TITLES_2),
    "matn:implications:single-hero": (("title", "item_1", "source"), ("item_1_title", *_HERO_PANEL)),
    "matn:limitations:numbered": (("title", *_ITEMS_4, *_markers(4), "source"), ()),
    "matn:closing:takeaways": (("title", *_ITEMS_3, "decisions_title", "decisions", "source"), _TITLES_3),
}
CHART_PLACEHOLDERS = {"chart"}
TITLE_PLACEHOLDERS = {"title"}

# Standard Windows/Office families; a face name passes when it is a family or "<family> <style>".
STANDARD_FONTS = frozenset({
    "aptos", "aptos display", "aptos narrow", "aptos serif", "aptos mono", "bahnschrift", "segoe ui", "calibri",
    "calibri light", "arial", "arial black", "arial narrow", "cambria", "candara", "corbel", "constantia",
    "consolas", "georgia", "tahoma", "verdana", "times new roman", "trebuchet ms", "franklin gothic", "gill sans mt",
    "century gothic", "courier new", "sitka", "symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings",
    "sakkal majalla", "traditional arabic", "simplified arabic", "arabic typesetting", "andalus", "aldhabi",
    "urdu typesetting", "microsoft uighur",
})
_STYLE_WORDS = {"light", "semilight", "semibold", "bold", "black", "italic", "condensed", "semicondensed",
                "display", "historic", "symbol", "emoji", "variable", "text", "small", "subheading", "heading",
                "medium", "demi", "book", "regular", "narrow"}


def is_standard_font(name: str) -> bool:
    words = name.lower().split()
    for cut in range(len(words), 0, -1):
        if " ".join(words[:cut]) in STANDARD_FONTS and all(w in _STYLE_WORDS for w in words[cut:]):
            return True
    return False


def layout_slide_type(name: str) -> tuple[str, str] | None:
    """("finding", "cards-3") for "matn:finding:cards-3"; None when the name is not a matn layout."""
    match = re.fullmatch(r"matn:([a-z_]+):([a-z0-9-]+)", name)
    return (match.group(1), match.group(2)) if match else None


@dataclass(frozen=True)
class TemplateReport:
    path: Path
    errors: tuple[str, ...]
    warnings: tuple[str, ...]

    @property
    def ok(self) -> bool:
        return not self.errors


def _fonts_in(xml: bytes) -> set[str]:
    root = etree.fromstring(xml)
    names = set()
    for tag in ("latin", "cs", "buFont"):
        for element in root.iter(f"{{{_NS['a']}}}{tag}"):
            face = element.get("typeface", "")
            if face and not face.startswith("+"):
                names.add(face)
    return names


def _theme_fonts(theme_xml: bytes) -> set[str]:
    root = etree.fromstring(theme_xml)
    names = set()
    for scheme in ("majorFont", "minorFont"):
        for tag in ("latin", "cs"):
            for element in root.iterfind(f".//a:fontScheme/a:{scheme}/a:{tag}", _NS):
                if element.get("typeface"):
                    names.add(element.get("typeface"))
    return names


def validate_template(path: Path) -> TemplateReport:
    """Check `path` against the template contract; errors make it unusable, warnings do not."""
    path = Path(path)
    errors: list[str] = []
    warnings: list[str] = []
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # any unreadable file is one clear error, not a traceback
        return TemplateReport(path, (f"TEMPLATE_UNREADABLE:{type(exc).__name__}",), ())

    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errors.append(f"SLIDE_SIZE_NOT_16_9:{prs.slide_width}x{prs.slide_height}")
    if len(prs.slides):
        errors.append(f"SAMPLE_SLIDES_LEFT:{len(prs.slides)}")

    known_types = {t.value for t in SlideType}
    seen: dict[str, int] = {}
    fonts: set[str] = set()
    for master in prs.slide_masters:
        fonts |= _fonts_in(master.part.blob)
        fonts |= _theme_fonts(master.part.part_related_by(RT.THEME).blob)
        for layout in master.slide_layouts:
            fonts |= _fonts_in(layout.part.blob)
            seen[layout.name] = seen.get(layout.name, 0) + 1
            parsed = layout_slide_type(layout.name)
            if parsed is None:
                warnings.append(f"LAYOUT_NOT_IN_CONTRACT:{layout.name}")
                continue
            if parsed[0] not in known_types:
                errors.append(f"LAYOUT_UNKNOWN_SLIDE_TYPE:{layout.name}")
                continue
            names = [shape.name for shape in layout.placeholders]
            for name in sorted({n for n in names if names.count(n) > 1}):
                errors.append(f"PLACEHOLDER_DUPLICATE:{layout.name}:{name}")
            by_name = {shape.name: shape for shape in layout.placeholders}
            required, optional = LAYOUTS.get(layout.name, (("title",), ()))
            for name in required:
                if name not in by_name:
                    errors.append(f"PLACEHOLDER_MISSING:{layout.name}:{name}")
            for name in (*required, *optional):
                shape = by_name.get(name)
                if shape is None:
                    continue
                kind = shape.placeholder_format.type
                if name in CHART_PLACEHOLDERS and kind != PP_PLACEHOLDER.CHART:
                    errors.append(f"PLACEHOLDER_WRONG_KIND:{layout.name}:{name}:{kind}")
                if name in TITLE_PLACEHOLDERS and kind not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    errors.append(f"PLACEHOLDER_WRONG_KIND:{layout.name}:{name}:{kind}")
    for name, count in seen.items():
        if count > 1 and layout_slide_type(name):
            errors.append(f"LAYOUT_DUPLICATE:{name}")
    for name in LAYOUTS:
        if name not in seen:
            errors.append(f"LAYOUT_MISSING:{name}")
    for face in sorted(fonts):
        if not is_standard_font(face):
            errors.append(f"FONT_NOT_STANDARD:{face}")
    return TemplateReport(path, tuple(errors), tuple(warnings))


@dataclass(frozen=True)
class GalleryTemplate:
    id: str
    path: Path
    thumbnail: Path
    meta: dict


def default_template_id(root: Path = GALLERY_DIR) -> str:
    """The owner's default gallery template (design/templates/src/gallery.json)."""
    return json.loads((root / "src" / "gallery.json").read_text(encoding="utf-8"))["default"]


def gallery(root: Path = GALLERY_DIR) -> tuple[GalleryTemplate, ...]:
    """Templates under `root`/<id>/ that have theme.json and pass the validator, sorted by id."""
    found = []
    for meta_path in sorted(root.glob("*/theme.json")):
        folder = meta_path.parent
        if validate_template(folder / "template.pptx").ok:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            found.append(GalleryTemplate(meta["id"], folder / "template.pptx", folder / "thumbnail.png", meta))
    return tuple(found)
