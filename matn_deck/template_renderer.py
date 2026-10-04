"""Renderer v2: DeckPlan + facts + tables -> editable .pptx through a gallery template, layout chosen by
generate-and-score (Step 2c).

The template owns fonts, colors, backgrounds, the title and the content area (docs/design/template-contract.md).
For each slide the renderer enumerates candidates: (variants valid for the content) x (type-scale steps) x
(anchors: top or centered). Each candidate's geometry comes from design/variants.json and measured text extents
(real font files), never from a PowerPoint export. Hard constraints reject a candidate (overflow, collision, bounds,
contrast, minimum sizes, empty slots, no visual on a content slide); soft scores pick the best one, with a
deterministic tie-break by variant order, step and anchor. When nothing fits, the slide splits; a fact is never cut.
Values (KPI numbers, chart data, pages, quotes, bar widths, icon counts) are copied or computed from facts and
tables by code.
"""
from __future__ import annotations

from dataclasses import dataclass, field, replace
from functools import lru_cache
import json
from pathlib import Path
import re
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

from matn_contracts.conversion import numeric_tokens
from matn_contracts.domain import ChartSpec, DeckPlan, Fact, Slide, SlideItem, Table, validate_deck_plan_against_facts
from matn_contracts.enums import FactKind, FactStatus, SlideType

from . import decor as decor_mod
from .chart_rules import ChartDataError, ChartDecision, decide, load_rules
from .colors import ColorResolver, blend, contrast, hex_rgb
from .decor import Decor
from .fit import WIDTH_SAFETY, Box, Fitter, TextStyle, placeholder_style
from .number_visuals import pictogram, progress_fraction
from .scoring import EMU_PER_PT, CandidateScore, Element, Rect, SlideContext, load_scoring, score
from .templates import LAYOUTS, validate_template
from .typography import balanced_lines, keep_caps
from .variants import (
    BLOCKS,
    gap_emu,
    has_inner_pad,
    load_variants,
    pad_edges,
    pad_emu,
    placeholder_name,
    slot_rects,
    sub_rect,
)

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_PAGE_IN_ID = re.compile(r":p(\d{4,}):")
_LEAD = re.compile(r"^(?P<lead>[^:]{1,40}):\s+(?P<rest>\S.*)$", re.S)
_YEAR = re.compile(r"(19|20)\d{2}")
_STEP = re.compile(r"(?:step|phase|stage)\s+(\d+)", re.I)
_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED_100,
}
ACCENT_RUNS_MAX = 2  # trick 16: numbers colored in one title
TITLE_STEPS = (1.0, 0.9, 0.8, 0.7)  # title slide and divider titles, never below the title minimum
VALUE_ROLES = frozenset({"kpi_value", "hero"})
CONTENT_TYPES = frozenset({SlideType.EXECUTIVE_SUMMARY, SlideType.KPI, SlideType.FINDING, SlideType.CHART,
                           SlideType.IMPLICATIONS, SlideType.LIMITATIONS, SlideType.CLOSING})
_FIXED_ROLES = {"body": "body", "quote": "quote", "decisions": "decisions", "decisions_title": "decisions_title",
                "attribution": "attribution"}
CHART_MIN_SHARE = 0.55


class DeckRenderError(ValueError):
    pass


class TemplateInvalidError(DeckRenderError):
    pass


@dataclass(frozen=True)
class Overflow:
    slide_number: int
    placeholder: str
    needed_emu: int
    available_emu: int


@dataclass(frozen=True)
class RenderReport:
    layouts: tuple[str, ...]  # layout used by each output slide, in order
    overflows: tuple[Overflow, ...]
    warnings: tuple[str, ...]
    scores: tuple[dict, ...] = field(default=(), compare=False)  # per output slide: chosen candidate and all scored


class DeckOverflowError(DeckRenderError):
    def __init__(self, report: RenderReport) -> None:
        where = ", ".join(f"slide {o.slide_number} {o.placeholder}" for o in report.overflows)
        super().__init__(f"TEXT_OVERFLOW: {where}")
        self.report = report


@lru_cache(maxsize=4)
def load_labels(language: str = "en") -> dict[str, Any]:
    return json.loads((Path(__file__).with_name("labels.json")).read_text(encoding="utf-8"))[language]


def page_number(block_or_table_id: str) -> int:
    """1-based physical PDF page encoded in a block_id or table_id."""
    match = _PAGE_IN_ID.search(block_or_table_id)
    if match is None:
        raise DeckRenderError(f"ID_HAS_NO_PAGE:{block_or_table_id}")
    return int(match.group(1)) + 1


@dataclass(frozen=True)
class _Content:
    """What one output slide shows. Splitting a slide splits `items`, `kpi_fact_ids` or `lines`."""
    slide_type: SlideType
    title: str
    items: tuple[SlideItem, ...] = ()
    kpi_fact_ids: tuple[str, ...] = ()
    chart: ChartSpec | None = None
    lines: tuple[str, ...] = ()
    line_fact_ids: tuple[str, ...] = ()
    notes: str | None = None
    section_number: int | None = None
    part: tuple[int, int] | None = None
    section: int | None = None  # the section (1-based divider count) a content slide belongs to
    plan_index: int = 0  # position of the plan slide this output slide comes from (split parts share it)

    @property
    def fact_ids(self) -> tuple[str, ...]:
        ids = [fid for item in self.items for fid in item.fact_ids] + list(self.kpi_fact_ids) + list(self.line_fact_ids)
        return tuple(dict.fromkeys(ids))


@dataclass
class _Plan:
    layout: Any
    fills: dict[str, list[str]]
    content: _Content
    overflows: list[Overflow] = field(default_factory=list)
    units: dict[str, str] = field(default_factory=dict)  # kpi_value_N -> unit run text
    step: str = "-"
    anchor: str = "top"
    sizes: dict[str, float] = field(default_factory=dict)  # placeholder -> font size set by code
    rects: dict[str, Rect] = field(default_factory=dict)  # placeholder -> box set by code
    anchors: dict[str, str] = field(default_factory=dict)  # placeholder -> vertical anchor set by code
    decor: list[Decor] = field(default_factory=list)
    elements: list[Element] = field(default_factory=list)
    score: CandidateScore | None = None
    order: tuple[int, int, int] = (0, 0, 0)  # variant, step, anchor index: the tie-break
    chart: ChartDecision | None = None
    body_pt: float | None = None  # the step's body size: KPI units and chart text use it (trick 17)
    title_accents: tuple[tuple[int, int], ...] = ()  # character spans of the title's accent numbers (trick 16)
    record_: dict = field(default_factory=dict)  # the scores.json entry: chosen candidate and every candidate

    @property
    def fits(self) -> bool:
        """A title overflow is reported but cannot be fixed by another variant or a split, so it does not count."""
        return all(o.placeholder == "title" for o in self.overflows)

    @property
    def ok(self) -> bool:
        return self.score.ok

    def record(self) -> dict:
        s = self.score
        return {"layout": self.layout.name, "step": self.step, "anchor": self.anchor,
                "total": None if s is None else s.total, "soft": None if s is None else s.soft,
                "hard": [] if s is None else s.hard}


def build_presentation(
    plan: DeckPlan,
    facts: tuple[Fact, ...],
    tables: tuple[Table, ...],
    template: Path,
    *,
    strict: bool = False,
    language: str = "en",
) -> tuple[Any, RenderReport]:
    report = validate_template(template)
    if not report.ok:
        raise TemplateInvalidError(f"TEMPLATE_INVALID:{Path(template).name}: " + "; ".join(report.errors))
    builder = _DeckBuilder(Presentation(str(template)), facts, tables, load_labels(language), Path(template))
    prs = builder.build(plan)
    result = RenderReport(tuple(builder.layouts_used), tuple(builder.overflows), tuple(builder.warnings),
                          tuple(builder.scores))
    if strict and result.overflows:
        raise DeckOverflowError(result)
    return prs, result


def render(
    plan: DeckPlan,
    facts: tuple[Fact, ...],
    tables: tuple[Table, ...],
    template: Path,
    out_path: Path,
    *,
    strict: bool = False,
    language: str = "en",
) -> RenderReport:
    prs, report = build_presentation(plan, facts, tables, template, strict=strict, language=language)
    prs.save(str(out_path))
    return report


def _rect(shape: Any) -> Rect:
    return Rect(shape.left, shape.top, shape.width, shape.height)


class _DeckBuilder:
    def __init__(self, prs: Any, facts: tuple[Fact, ...], tables: tuple[Table, ...], labels: dict[str, Any],
                 template: Path) -> None:
        self.prs = prs
        self.facts = {fact.fact_id: fact for fact in facts}
        self.tables = {table.table_id: table for table in tables}
        self.labels = labels
        self.fitter = Fitter()
        self.layouts = {layout.name: layout for layout in prs.slide_layouts}
        self.layouts_used: list[str] = []
        self.overflows: list[Overflow] = []
        self.warnings: list[str] = []
        self.scores: list[dict] = []
        self._styles: dict[tuple[str, str], TextStyle] = {}
        self.cfg = load_scoring()
        self.variants = load_variants()["variants"]
        self.geo = load_variants()["geometry"]
        meta_path = template.with_name("theme.json")
        self.meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.is_file() else {}
        self.card_style = {**decor_mod.DEFAULT_CARD_STYLE, **self.meta.get("card_style", {})}
        self.card_style["shadow"] = {**decor_mod.DEFAULT_CARD_STYLE["shadow"],
                                     **self.card_style.get("shadow", {})}
        self.divider_style = self.meta.get("divider_style", {"type": "color_block"})
        reference = self.layouts["matn:limitations:default"]
        ref = {ph.name: ph for ph in reference.placeholders}
        self.area = _rect(ref["body"])  # the contract's content area (template-contract.md §5)
        self.title_box = _rect(ref["title"])
        chip = _rect(ref["source"])
        self.chip_band = Rect(self.area.x, self.area.bottom, self.area.w, chip.bottom - self.area.bottom)
        self.chip_h = chip.h
        self._current_kpis: tuple[str, ...] = ()
        self.slide = Rect(0, 0, prs.slide_width, prs.slide_height)
        self.colors = ColorResolver(prs.slide_masters[0])
        accent = next(ph for ph in self.layouts["matn:finding:cards-3"].placeholders if ph.name == "item_1_title")
        self.accent_xml = self._style(self.layouts["matn:finding:cards-3"], accent).color_xml
        if self.accent_xml is None:
            raise TemplateInvalidError("TEMPLATE_INVALID: item titles carry no accent color")
        self.accent_xml = etree.tostring(etree.fromstring(self.accent_xml)[0]).decode()  # the color inside solidFill
        self.focus_xml = (decor_mod.tint("accent2", 1.0)
                          if self.meta.get("family") == "creative" else self.accent_xml)
        self.title_pt = float(self.cfg["title_sizes"][-1])
        # Trick 17: the template's three content sizes (plus its hero display), read from its own layouts.
        self.type_sizes = {
            "display": self._layout_size("matn:kpi:cards", "kpi_value_1"),
            "hero": self._layout_size("matn:kpi:hero-number", "kpi_value_1"),
            "body": self._layout_size("matn:finding:cards-3", "item_1"),
        }

    def _layout_size(self, layout: str, placeholder: str) -> float:
        return self._style(self.layouts[layout], self._layout_ph(self.layouts[layout], placeholder)).size_pt

    # ---- deck ---------------------------------------------------------------------------------------------------

    def build(self, plan: DeckPlan) -> Any:
        validate_deck_plan_against_facts(plan, tuple(self.facts.values()))
        for slide in plan.slides:
            if slide.chart is not None and slide.chart.table_id not in self.tables:
                raise DeckRenderError(f"UNKNOWN_TABLE_ID:{slide.chart.table_id}")
            for fact_id in self._slide_fact_ids(slide):
                status = self.facts[fact_id].status
                if status is FactStatus.REJECTED:
                    raise DeckRenderError(f"REJECTED_FACT_CITED:{fact_id}")
                if status is not FactStatus.VERIFIED and f"UNVERIFIED_FACT:{fact_id}" not in self.warnings:
                    self.warnings.append(f"UNVERIFIED_FACT:{fact_id}")
        self.prs.core_properties.title = plan.slides[0].title
        self.title_pt = self._deck_title_size(plan)
        self.sections = sum(s.slide_type is SlideType.SECTION_DIVIDER for s in plan.slides)

        cited = list(dict.fromkeys(fid for slide in plan.slides for fid in self._slide_fact_ids(slide)))
        contents = []
        dividers = 0
        for index, slide in enumerate(plan.slides):
            content = _Content(slide.slide_type, slide.title, slide.items, slide.kpi_fact_ids, slide.chart,
                               notes=slide.notes, section=dividers or None, plan_index=index)
            if slide.slide_type is SlideType.SECTION_DIVIDER:
                dividers += 1
                content = replace(content, section_number=dividers)
            elif slide.slide_type is SlideType.EVIDENCE_APPENDIX:
                content = self._appendix(content, cited)
            contents.append(content)

        planned: list[list[_Plan]] = []
        previous: str | None = None
        for content in contents:
            parts = self._plan(content, previous)
            planned.append(parts)
            last = parts[-1]
            previous = last.layout.name if content.slide_type in CONTENT_TYPES else None
        for index, slide in enumerate(plan.slides):
            if slide.slide_type is SlideType.SOURCES:
                lines = self._sources_lines(plan, planned, cited)
                planned[index] = self._plan(replace(contents[index], lines=tuple(lines)), None)
        for parts in planned:
            for part in parts:
                self._emit(part)
        return self.prs

    def _slide_fact_ids(self, slide: Slide) -> list[str]:
        return list(dict.fromkeys([fid for item in slide.items for fid in item.fact_ids] + list(slide.kpi_fact_ids)))

    def _deck_title_size(self, plan: DeckPlan) -> float:
        """One title size for every content slide: the largest of `title_sizes` at which every title fits."""
        layout = self.layouts["matn:limitations:default"]
        title = next(ph for ph in layout.placeholders if ph.name == "title")
        titles = [s.title for s in plan.slides if s.slide_type in CONTENT_TYPES]
        for size in self.cfg["title_sizes"]:
            if all(self._extent([t], self._title_style(layout, title, size, t), title.width)[0] <= title.height
                   for t in titles):
                return float(size)
        return float(self.cfg["title_sizes"][-1])

    # ---- candidates ---------------------------------------------------------------------------------------------

    def _variant_order(self, slide_type: SlideType) -> list[str]:
        prefix = f"matn:{slide_type.value}:"
        order = [name for name in LAYOUTS if name.startswith(prefix)]
        order += sorted(n for n in self.layouts if n.startswith(prefix) and n not in order)
        return sorted(order, key=lambda n: not self.variants.get(n, {}).get("specific"))  # stable: specific first

    def _valid(self, name: str, content: _Content) -> bool:
        """True when the variant can show all of `content` (nothing dropped) and its data rules hold."""
        if name not in self.layouts:
            return False
        spec = self.variants.get(name)
        if spec is None:
            return False
        names = {ph.name for ph in self.layouts[name].placeholders}
        valid = spec.get("valid", {})
        n_items, n_kpis = len(content.items), len(content.kpi_fact_ids)
        if content.lines:
            return "body" in names
        if content.slide_type in (SlideType.TITLE, SlideType.SECTION_DIVIDER):
            return not content.items or "subtitle" in names
        takeaways, decisions = self._split_decisions(content)
        if "takeaways" in valid:
            low, high = valid["takeaways"]
            if not low <= len(takeaways) <= high:
                return False
        elif "items" in valid:
            low, high = valid["items"]
            if not low <= n_items <= high:
                return False
        elif n_items and "body" not in names:
            return False
        if "kpis" in valid:
            low, high = valid["kpis"]
            if not low <= n_kpis <= high:
                return False
        elif n_kpis:
            return False
        if n_items and "kpis" in valid and "body" not in names:
            return False
        values = [self.facts[fid].number_value for fid in content.kpi_fact_ids]
        rule = valid.get("requires")
        if rule == "pictogram":
            return all(pictogram(v) is not None for v in values)
        if rule == "progress":
            return all(progress_fraction(v) is not None for v in values)
        if rule == "leads":
            return all(_LEAD.match(item.text) for item in content.items)
        if rule == "ordered":
            return self._ordered(content.items)
        if rule == "quote":
            return bool(content.items)
        if rule == "decisions":
            return bool(decisions) and bool(takeaways)
        return True

    @staticmethod
    def _ordered(items: tuple[SlideItem, ...]) -> bool:
        """Items whose leads are increasing years or numbered steps (a timeline)."""
        leads = [_LEAD.match(item.text) for item in items]
        if not all(leads):
            return False
        texts = [lead["lead"].strip() for lead in leads]
        for pattern, key in ((_YEAR, lambda m: int(m.group(0))), (_STEP, lambda m: int(m.group(1)))):
            matches = [pattern.fullmatch(t) for t in texts]
            if all(matches):
                keys = [key(m) for m in matches]
                return all(a < b for a, b in zip(keys, keys[1:]))
        return False

    def _split_decisions(self, content: _Content) -> tuple[list[SlideItem], list[SlideItem]]:
        """Closing slides: items citing only recommendation facts are decisions needed; the rest are takeaways."""
        if content.slide_type is not SlideType.CLOSING:
            return list(content.items), []
        takeaways, decisions = [], []
        for item in content.items:
            kinds = {self.facts[fid].kind for fid in item.fact_ids}
            (decisions if kinds == {FactKind.RECOMMENDATION} else takeaways).append(item)
        return takeaways, decisions

    def _plan(self, content: _Content, previous: str | None) -> list[_Plan]:
        attempts = self._attempts(content, previous)
        good = [a for a in attempts if a.ok and a.fits]
        if good:
            varied = [a for a in good if a.layout.name != previous]
            if varied:
                good = varied  # deck variety is a contract invariant when another valid layout exists
            top = max(self._preference(a) for a in good)
            tied = [a for a in good if self._preference(a) >= top - self.cfg["tie_margin"]]
            return [self._remember(min(tied, key=lambda a: a.order), attempts)]
        halves = self._split(content)
        if halves is None:
            # Nothing passes and nothing can split: keep the least-bad candidate and report its violations.
            chosen = min(attempts, key=lambda a: (not a.fits, len(a.score.hard), -a.score.total, a.order))
            return [self._remember(chosen, attempts)]
        parts = self._plan(halves[0], previous)
        parts += self._plan(halves[1], parts[-1].layout.name if content.slide_type in CONTENT_TYPES else None)
        if content.part is not None:
            return parts
        renumbered, before = [], previous
        for k, part in enumerate(parts, start=1):
            redo = self._evaluate(part.layout, replace(part.content, part=(k, len(parts))), part.order, before)
            redo.record_ = {**part.record_, "chosen": redo.record()}
            renumbered.append(redo)
            before = redo.layout.name if content.slide_type in CONTENT_TYPES else None
        return renumbered

    def _preference(self, attempt: _Plan) -> float:
        """The total used to choose among candidates that pass every hard constraint: the score plus the template's
        small layout-bias bonus (theme.json "layout_bias": variant names) for its favourite variants. The recorded
        score, and so the beauty gate, stays the unbiased total; a bias never rescues a failing candidate."""
        variant = attempt.layout.name.rsplit(":", 1)[-1]
        bonus = self.cfg.get("layout_bias_bonus", 0.0) if variant in self.meta.get("layout_bias", ()) else 0.0
        return attempt.score.total + bonus

    @staticmethod
    def _remember(chosen: _Plan, attempts: list[_Plan]) -> _Plan:
        chosen.record_ = {"slide_type": chosen.content.slide_type.value, "chosen": chosen.record(),
                          "candidates": [a.record() for a in sorted(attempts, key=lambda a: a.order)]}
        return chosen

    def _attempts(self, content: _Content, previous: str | None) -> list[_Plan]:
        names = [n for n in self._variant_order(content.slide_type) if self._valid(n, content)]
        if not names:
            raise DeckRenderError(f"NO_LAYOUT_FOR:{content.slide_type.value}:{len(content.items)} items, "
                                  f"{len(content.kpi_fact_ids)} kpis")
        scored = content.slide_type in CONTENT_TYPES
        steps = range(len(self.cfg["type_scale"])) if scored else [0]
        anchors = range(len(self.cfg["anchors"])) if scored else [0]
        attempts = []
        for v, name in enumerate(names):
            for s in steps:
                for a in anchors:
                    attempts.append(self._evaluate(self.layouts[name], content, (v, s, a), previous))
        return attempts

    def _split(self, content: _Content) -> tuple[_Content, _Content] | None:
        if len(content.lines) >= 2:
            mid = (len(content.lines) + 1) // 2
            first_ids, second_ids = content.line_fact_ids[:mid], content.line_fact_ids[mid:]
            return (replace(content, lines=content.lines[:mid], line_fact_ids=first_ids),
                    replace(content, lines=content.lines[mid:], line_fact_ids=second_ids))
        if len(content.kpi_fact_ids) >= 2 and content.slide_type is SlideType.KPI:
            mid = (len(content.kpi_fact_ids) + 1) // 2
            return (replace(content, kpi_fact_ids=content.kpi_fact_ids[:mid]),
                    replace(content, kpi_fact_ids=content.kpi_fact_ids[mid:], items=()))
        if len(content.items) >= 2 and content.slide_type not in (SlideType.CHART, SlideType.TITLE):
            mid = (len(content.items) + 1) // 2
            return replace(content, items=content.items[:mid]), replace(content, items=content.items[mid:])
        return None

    # ---- one candidate: geometry, measurement, score --------------------------------------------------------------

    def _style(self, layout: Any, ph: Any) -> TextStyle:
        key = (layout.name, ph.name)
        if key not in self._styles:
            self._styles[key] = placeholder_style(ph)
        return self._styles[key]

    def _title_style(self, layout: Any, ph: Any, size: float, text: str) -> TextStyle:
        style = self._style(layout, ph)
        return replace(style, size_pt=size, caps=keep_caps(text, style.caps))

    def _extent(self, paragraphs: list[str], style: TextStyle, width: float) -> tuple[float, float]:
        """(height incl. insets, widest line incl. insets) of `paragraphs` wrapped at `width`."""
        inner = (width - style.insets[0] - style.insets[2] - style.indent_emu) * WIDTH_SAFETY
        widest = 0.0
        for text in paragraphs:
            for line in self.fitter.wrap(text, inner, style):
                widest = max(widest, self.fitter.width(line, style))
        height = self.fitter.needed(paragraphs, style, Box(int(width), 0))
        return height, min(width, widest / WIDTH_SAFETY + style.insets[0] + style.insets[2] + style.indent_emu)

    def _fg(self, style: TextStyle) -> tuple[float, float, float]:
        if style.color_xml:
            return self.colors.resolve(style.color_xml)[0]
        return self.colors.scheme("tx1")

    def _contrast(self, style: TextStyle, layout: Any, backdrop: str | None) -> float:
        backgrounds = self.colors.background(layout)
        if backdrop is not None:
            alpha = float(self.card_style["fill_alpha"])
            top = self.colors.scheme("tx1")
            backgrounds = [blend(top, b, alpha) for b in backgrounds]
        fg = self._fg(style)
        return min(contrast(fg, b) for b in backgrounds)


    def _evaluate(self, layout: Any, content: _Content, order: tuple[int, int, int], previous_variant: str | None
                  ) -> _Plan:
        spec = self.variants.get(layout.name, {})
        scored = content.slide_type in CONTENT_TYPES
        step = self.cfg["type_scale"][order[1]] if scored else {}
        anchor = self.cfg["anchors"][order[2]]
        self._current_kpis = content.kpi_fact_ids
        fills = self._fills(layout, content, spec)
        units = {f"kpi_value_{n}": self._unit(fid) for n, fid in enumerate(content.kpi_fact_ids, start=1)}
        hero = self._hero_number(content) if spec.get("hero_panel") else None
        if hero is not None:
            units["kpi_value_1"] = self._unit(hero.fact_id)
        plan = _Plan(layout, fills, content, units={k: v for k, v in units.items() if v and k in fills},
                     step=step.get("name", "-"), anchor=anchor, order=order)
        m = _Measure(self, plan, step, scored)
        plan.body_pt = m.size_for("body") if scored else None
        elements: list[Element] = [self._title_element(m, scored)]

        fixed = spec.get("fixed", {})
        region_frac = spec.get("slots", {}).get("region")
        region_bottom = region_frac[1] + region_frac[3] if region_frac else None
        below = [n for n, f in fixed.items() if region_bottom is not None and f[1] >= region_bottom - 1e-9]
        cards = [(sub_rect(self.area, c), c) for c in spec.get("fixed_cards", [])]
        below_cards = [r for r, c in cards if region_bottom is not None and c[1] >= region_bottom - 1e-9]
        side_cards = [r for r, c in cards if r not in below_cards]

        # Side placeholders: a card hugs what it holds; free text keeps its rect with the candidate's anchor.
        in_card: set[str] = set()
        for k, card in enumerate(side_cards, start=1):
            names = sorted((n for n, f in fixed.items() if n in fills and card.contains(sub_rect(self.area, f), 1)),
                           key=lambda n: fixed[n][1])
            in_card |= set(names)
            elements += self._hug_fixed(m, spec, card, names, anchor, f"card_f{k}")
        for name, frac in fixed.items():
            if name in below or name in in_card:
                continue
            rect = sub_rect(self.area, frac)
            if name == "chart":
                if content.chart is not None:
                    plan.rects["chart"] = rect
                    elements.append(Element("chart", "chart", rect, ink=0.6))
                    plan.chart = self._decide_chart(content)
                continue
            if name in fills:
                role = None if spec.get("list") or not scored else _FIXED_ROLES.get(name)
                elements.append(m.place(name, rect, role, None, anchor, "quote" if name == "quote" else "text"))

        panel = spec.get("hero_panel")
        if panel:
            region = sub_rect(self.area, panel["region"])
            names = [n for n in panel["number"] if n in fills] or [n for n in panel["quote"] if n in fills]
            roles = {"kpi_value_1": "hero", "kpi_label_1": "kpi_label", "quote": "quote", "attribution": "attribution"}
            kinds = {"kpi_value_1": "hero", "quote": "quote"}
            parts = [_Part(n, roles[n], kinds.get(n, "text")) for n in names]
            pad = pad_emu(spec)
            m.fit_values([(parts, region.inset(pad))])
            need = m.height_of(parts, region.inset(pad)) + 2 * pad
            top = region.y if anchor == "top" else region.y + max(0.0, (region.h - need) / 2)
            card = Rect(region.x, top, region.w, min(need, region.h))
            plan.decor.append(Decor("hero_card", card, "hero_card"))
            elements.append(Element("hero_card", "card", card, container=True, empty=not names,
                                    overflow=need > region.h + EMU_PER_PT))
            elements += m.stack(parts, card.inset(pad), 0.0, "hero_card")

        if spec.get("slots"):
            height, place_slots = self._slot_group(m, spec, content)
            below_h, place_below = self._below_block(m, spec, [n for n in below if n in fills], below_cards)
            gap = gap_emu(spec)
            total = height + (gap + below_h if below_h else 0.0)
            top = self.area.y + region_frac[1] * self.area.h
            room = self.area.bottom - top
            start = top if anchor == "top" else top + max(0.0, (room - total) / 2)
            if total > room + EMU_PER_PT:
                plan.overflows.append(Overflow(0, "slots", int(total), int(room)))
                elements.append(Element("slots", "text", Rect(self.area.x, top, self.area.w, room), overflow=True,
                                        container=True))
            elements += place_slots(start)
            if below_h:
                elements += place_below(start + height + gap)

        if "source" in fills:
            src = m.by_name["source"]
            width = max(self.fitter.text_width(fills["source"][0], self._style(layout, src)), src.height * 2)
            elements.append(Element("source", "chip", Rect(src.left, src.top, width, src.height), ink=0.8))
        if scored and content.section and self.sections >= 2:
            dots = self._tracker(content.section)
            plan.decor += dots
            elements.append(Element("tracker", "tracker", Rect(dots[0].rect.x, dots[0].rect.y,
                                                               dots[-1].rect.right - dots[0].rect.x, dots[0].rect.h)))
        if content.slide_type is SlideType.SECTION_DIVIDER:
            for name in ("section_number", "subtitle"):
                if name in fills:
                    elements.append(m.place(name, _rect(m.by_name[name]), None, None, "top"))
        if content.slide_type is SlideType.TITLE and "subtitle" in fills:
            subtitle = m.place("subtitle", _rect(m.by_name["subtitle"]), None, None, "top")
            elements.append(replace(subtitle, contrast=None))  # on the cover art: measured by the gallery QA

        plan.elements = elements
        divider = content.slide_type is SlideType.SECTION_DIVIDER
        ctx = SlideContext(content_slide=scored, area=self.area,
                           safe=self._divider_safe() if divider else self.area if scored else self.slide,
                           slide=self.slide, title_pt=self.title_pt if scored else None,
                           centered_hero=bool(spec.get("hero") or panel) and anchor == "center")
        variety = 0.0 if previous_variant == layout.name else 1.0
        plan.score = score(elements, ctx, variety=variety, cfg=self.cfg)
        chart = plan.rects.get("chart")
        if scored and chart is not None and chart.area / self.area.area < CHART_MIN_SHARE:
            plan.score.hard.append("CHART_TOO_SMALL")
        if scored:
            sizes = self.content_sizes(plan)
            if len(sizes) > self.cfg["max_content_sizes"]:
                plan.score.hard.append(f"TOO_MANY_SIZES:{len(sizes)}")
        return plan

    def content_sizes(self, plan: _Plan) -> set[float]:
        """Distinct text sizes a content slide uses (trick 17): title, placeholders, KPI units and chart text."""
        sizes = {float(v) for k, v in plan.sizes.items()}
        if plan.units or plan.chart is not None:
            sizes.add(float(plan.body_pt))
        return sizes

    def _title_element(self, m: "_Measure", scored: bool) -> Element:
        layout, plan = m.layout, m.plan
        ph = m.by_name["title"]
        text = plan.fills["title"][0]
        if scored:
            size = self.title_pt
        else:  # title slide and dividers: the template's size, stepped down until the title fits
            base = self._style(layout, ph).size_pt
            sizes = [max(self.cfg["minimums"]["title_pt"], round(base * f)) for f in TITLE_STEPS]
            size = next((s for s in sizes if self._extent([text], self._title_style(layout, ph, s, text),
                                                          ph.width)[0] <= ph.height), sizes[-1])
        style = self._title_style(layout, ph, size, text)
        needed, width = self._extent([text], style, ph.width)
        if needed > ph.height:
            plan.overflows.append(Overflow(0, "title", int(needed), ph.height))
        if scored or size != self._style(layout, ph).size_pt:
            plan.sizes["title"] = size
        # The template draws the title over its own art (bands, photos), so only a divider title, whose full-bleed
        # fill is known, and accent numbers on a title that sits on the slide background are checked here; the
        # gallery QA measures every title from the exported pixels.
        ratio = None
        if plan.content.slide_type is SlideType.SECTION_DIVIDER:
            ratio = self._contrast(style, layout, None)
        elif scored:
            plan.title_accents = self._title_accents(plan.content, text, layout, style)
            if plan.title_accents:
                ratio = min(self._contrast(style, layout, None), self._text_accent_contrast(None))
        return Element("title", "title", Rect(ph.left, ph.top, width, ph.height), ink=0.6, size_pt=size,
                       bold=style.bold, contrast=ratio)

    def _title_accents(self, content: _Content, text: str, layout: Any, style: TextStyle) -> tuple[tuple[int, int], ...]:
        """Trick 16: character spans of at most two numbers in an assertion title, found by the contracts' number
        tokenizer; a number equal to a cited fact's value comes first. Only on titles that sit on the slide
        background (title and accent both readable there), so the text contrast check still holds."""
        if self._contrast(style, layout, None) < self.cfg["contrast"]["large"] or \
                self._text_accent_contrast(None) < self.cfg["contrast"]["large"]:
            return ()
        cited = {self.facts[f].number_value.value_as_written.strip() for f in content.fact_ids
                 if self.facts[f].number_value}
        spans = []
        for match in numeric_tokens(text):
            end = match.end() + (1 if text[match.end():match.end() + 1] == "%" else 0)
            spans.append((match.group(0) not in cited, match.start(), end))
        chosen = sorted(spans)[:ACCENT_RUNS_MAX]
        return tuple(sorted((start, end) for _, start, end in chosen))

    def _slot_group(self, m: "_Measure", spec: dict, content: _Content):
        """Measure the slots, hug each card to its content (cards in one row share the tallest height) and return
        (group height, place(top) -> elements)."""
        slots = spec["slots"]
        stems = slots.get("stack", []) + slots.get("left", []) + slots.get("right", [])
        if "kpi_value" in stems:
            count = len(content.kpi_fact_ids)
        else:
            count = len(self._split_decisions(content)[0]) if "decisions" in m.plan.fills else len(content.items)
        region = sub_rect(self.area, slots["region"])
        gap = gap_emu(spec)
        rects = slot_rects(slots["recipe"], region, count, gap)
        pad = pad_emu(spec) if has_inner_pad(spec) else 0.0
        roles = spec.get("roles", {})
        side = slots.get("arrange") == "side"
        half_gap = self.geo["gap_pt"] * EMU_PER_PT / 2
        entries = []  # (slot, parts, left parts, left width)
        for k, slot in enumerate(rects, start=1):
            inner_w = slot.w - 2 * pad
            if side:
                left_w = self.geo["marker_pt"] * EMU_PER_PT if slots["split"] == "marker" else inner_w * slots["split"]
                entries.append((slot, self._parts(m, slots["right"], k, roles), self._parts(m, slots["left"], k, roles),
                                left_w))
            else:
                entries.append((slot, self._parts(m, slots["stack"], k, roles), [], 0.0))

        def columns(slot: Rect, left_w: float) -> tuple[Rect, Rect | None]:
            inner = slot.inset(pad) if pad else slot
            if not side:
                return inner, None
            return (Rect(inner.x + left_w + half_gap, inner.y, inner.w - left_w - half_gap, inner.h),
                    Rect(inner.x, inner.y, left_w, inner.h))

        m.fit_values([(parts, columns(slot, lw)[0]) for slot, parts, _, lw in entries])
        needs = []
        for slot, parts, left, lw in entries:
            column, left_col = columns(slot, lw)
            need = m.height_of(parts, column)
            if left:
                need = max(need, m.height_of(left, left_col))
            needs.append(need + 2 * pad)
        # Hug: rows of the recipe become bands as tall as their tallest single-band card.
        bands = sorted({round(slot.y) for slot, *_ in entries})
        covers = [[b for b, y in enumerate(bands) if slot.y - 1 <= y < slot.bottom - 1] for slot, *_ in entries]
        band_h = [0.0] * len(bands)
        for cover, need in zip(covers, needs):
            if len(cover) == 1:
                band_h[cover[0]] = max(band_h[cover[0]], need)
        for cover, need in zip(covers, needs):
            if len(cover) > 1:
                spanned = sum(band_h[b] for b in cover) + gap * (len(cover) - 1)
                band_h[cover[-1]] += max(0.0, need - spanned)
        offsets = [sum(band_h[:b]) + gap * b for b in range(len(bands))]
        height = sum(band_h) + gap * (len(bands) - 1)

        def place(top: float) -> list[Element]:
            elements: list[Element] = []
            backdrop = "card" if spec.get("cards") else None
            for k, ((slot, parts, left, lw), cover) in enumerate(zip(entries, covers), start=1):
                h = sum(band_h[b] for b in cover) + gap * (len(cover) - 1)
                card = Rect(slot.x, top + offsets[cover[0]], slot.w, h)
                if spec.get("cards"):
                    m.plan.decor.append(Decor("card", card, f"card_{k}"))
                    elements.append(Element(f"card_{k}", "card", card, container=True, empty=not (parts or left)))
                    if spec.get("strip"):
                        inset = 10 * EMU_PER_PT
                        m.plan.decor.append(Decor("strip", Rect(card.x + inset, card.y, card.w - 2 * inset,
                                                                self.geo["strip_pt"] * EMU_PER_PT), f"strip_{k}"))
                column, left_col = columns(card, lw)
                content_h = max(m.height_of(parts, column), m.height_of(left, left_col) if left else 0.0)
                offset = 0.0 if m.plan.anchor == "top" else max(0.0, (column.h - content_h) / 2)
                elements += m.stack(parts, column, offset, backdrop)
                if left:
                    elements += m.stack(left, left_col, offset, backdrop)
            dots = [d for d in m.plan.decor if d.kind == "timeline_dot"]
            if spec.get("timeline") and len(dots) >= 2:
                cy = dots[0].rect.y + dots[0].rect.h / 2
                thick = 3 * EMU_PER_PT
                x0, x1 = dots[0].rect.x + dots[0].rect.w / 2, dots[-1].rect.x + dots[-1].rect.w / 2
                m.plan.decor.append(Decor("timeline_line", Rect(x0, cy - thick / 2, x1 - x0, thick), "timeline_line"))
                box = Rect(dots[0].rect.x, min(d.rect.y for d in dots), dots[-1].rect.right - dots[0].rect.x,
                           dots[0].rect.h)
                elements.append(Element("timeline", "timeline", box, ink=0.7, contrast=self._graphic_contrast(None)))
            return elements

        return height, place

    def _below_block(self, m: "_Measure", spec: dict, names: list[str], cards: list[Rect]):
        """Placeholders under the slot group (a KPI slide's line, the closing decisions card): measured now, placed
        right under the group so the whole composition anchors as one."""
        if not names:
            return 0.0, None
        fixed = spec["fixed"]
        card = cards[0] if cards else None
        pad = pad_emu(spec) if card else 0.0
        rects = {}
        for name in names:
            rect = sub_rect(self.area, fixed[name])
            rects[name] = pad_edges(rect, card, pad) if card else rect
        roles = {n: (None if spec.get("list") else _FIXED_ROLES.get(n)) for n in names}
        heights = {n: m.part_height(_Part(n, roles[n], "text"), rects[n].w) for n in names}
        block = max(heights.values()) + 2 * pad

        def place(top: float) -> list[Element]:
            elements: list[Element] = []
            if card is not None:
                box = Rect(card.x, top, card.w, block)
                m.plan.decor.append(Decor("card", box, "card_below"))
                elements.append(Element("card_below", "card", box, container=True))
            for name in names:
                r = rects[name]
                elements.append(m.place(name, Rect(r.x, top + pad, r.w, heights[name] + self.geo["slack_pt"] * EMU_PER_PT),
                                        roles[name], "card" if card else None, "top"))
            return elements

        return block, place

    def _hug_fixed(self, m: "_Measure", spec: dict, card: Rect, names: list[str], anchor: str, label: str
                   ) -> list[Element]:
        """A fixed card (the split slide's quote card) hugs its stacked placeholders and anchors in its own rect."""
        pad = pad_emu(spec)
        gap = self.geo["stack_gap_pt"] * EMU_PER_PT
        roles = {n: _FIXED_ROLES.get(n) for n in names}
        heights = [m.part_height(_Part(n, roles[n], "text"), card.w - 2 * pad) for n in names]
        need = sum(heights) + gap * max(0, len(names) - 1) + 2 * pad
        top = card.y if anchor == "top" else card.y + max(0.0, (card.h - need) / 2)
        box = Rect(card.x, top, card.w, min(need, card.h))
        m.plan.decor.append(Decor("card", box, label))
        elements = [Element(label, "card", box, container=True, empty=not names, overflow=need > card.h + EMU_PER_PT)]
        y = top + pad
        slack = self.geo["slack_pt"] * EMU_PER_PT
        for name, h in zip(names, heights):
            kind = "quote" if name == "quote" else "text"
            elements.append(m.place(name, Rect(card.x + pad, y, card.w - 2 * pad, h + slack), roles[name], "card", "top",
                                    kind))
            y += h + gap
        return elements

    @staticmethod
    def _parts(m: "_Measure", stems: list[str], k: int, roles: dict) -> list["_Part"]:
        out = []
        for stem in stems:
            if stem in BLOCKS:
                out.append(_Part(stem, None, "block", k))
                continue
            name = placeholder_name(stem, k)
            if name not in m.plan.fills:
                continue
            role = roles.get(stem, stem)
            kind = {"kpi_value": "hero" if role == "hero" else "kpi", "marker": "marker"}.get(stem, "text")
            out.append(_Part(name, role, kind, k))
        return out

    def _graphic_contrast(self, backdrop: str | None) -> float:
        return self._accent_contrast(self.focus_xml, backdrop)

    def _text_accent_contrast(self, backdrop: str | None) -> float:
        return self._accent_contrast(self.accent_xml, backdrop)

    def _accent_contrast(self, color_xml: str, backdrop: str | None) -> float:
        accent = self.colors.resolve(color_xml)[0]
        backgrounds = self.colors.background(None)
        if backdrop == "card":
            backgrounds = [blend(self.colors.scheme("tx1"), b, float(self.card_style["fill_alpha"]))
                           for b in backgrounds]
        return min(contrast(accent, b) for b in backgrounds)

    def _block_height(self, part: "_Part", width: float) -> float:
        g = self.geo
        if part.name == "dot":
            return g["dot_pt"] * EMU_PER_PT
        if part.name == "bar":
            return g["bar_pt"] * EMU_PER_PT
        counts = pictogram(self.facts[self._current_kpis[part.slot - 1]].number_value)
        return self._icon_layout(counts[1], width)[2] if counts else 0.0

    def _icon_layout(self, total: int, width: float) -> tuple[float, int, float]:
        """(icon size, icons per row, block height) for `total` icons in `width`."""
        g = self.geo
        gap = g["icon_gap_pt"] * EMU_PER_PT
        per_row, rows = total, 1
        size = min(g["icon_max_pt"] * EMU_PER_PT, (width - (per_row - 1) * gap) / per_row)
        if size < 14 * EMU_PER_PT:
            per_row, rows = (total + 1) // 2, 2
            size = min(g["icon_max_pt"] * EMU_PER_PT, (width - (per_row - 1) * gap) / per_row)
        return size, per_row, rows * size + (rows - 1) * gap

    def _place_block(self, m: "_Measure", part: "_Part", rect: Rect, backdrop: str | None) -> list[Element]:
        k = part.slot
        if part.name == "dot":
            m.plan.decor.append(Decor("timeline_dot", Rect(rect.x, rect.y, rect.h, rect.h), f"dot_{k}"))
            return []
        fact = self.facts[self._current_kpis[k - 1]]
        if part.name == "bar":  # trick 15: width = value / 100 of the track
            fraction = progress_fraction(fact.number_value)
            m.plan.decor.append(Decor("bar_track", rect, f"bar_track_{k}"))
            m.plan.decor.append(Decor("bar_fill", Rect(rect.x, rect.y, round(rect.w * float(fraction)), rect.h),
                                      f"bar_fill_{k}"))
            return [Element(f"bar_{k}", "bar", rect, ink=0.8, contrast=self._graphic_contrast(backdrop))]
        filled, total = pictogram(fact.number_value)  # trick 14: icon counts from the fact's own numbers
        size, per_row, _ = self._icon_layout(total, rect.w)
        gap = self.geo["icon_gap_pt"] * EMU_PER_PT
        for i in range(total):
            x = rect.x + (i % per_row) * (size + gap)
            y = rect.y + (i // per_row) * (size + gap)
            m.plan.decor.append(Decor("icon_on" if i < filled else "icon_off", Rect(x, y, size, size),
                                      f"icon_{k}_{i + 1}"))
        box = Rect(rect.x, rect.y, min(rect.w, per_row * (size + gap) - gap), rect.h)
        return [Element(f"pictogram_{k}", "icon", box, ink=0.8, contrast=self._graphic_contrast(backdrop))]

    def _divider_safe(self) -> Rect:
        frac = (self.meta.get("divider") or {}).get("safe")
        return sub_rect(self.slide, frac) if frac else self.slide

    def _tracker(self, section: int) -> list[Decor]:
        """Trick 3: one dot per section in the chip band's right end; the current section's dot in the accent."""
        size = self.geo["tracker_dot_pt"] * EMU_PER_PT
        gap = self.geo["tracker_gap_pt"] * EMU_PER_PT
        n = self.sections
        band = self.chip_band
        x0 = band.right - (n * size + (n - 1) * gap)
        y = band.bottom - self.chip_h / 2 - size / 2  # level with the chip's center
        return [Decor("tracker_on" if k == section else "tracker_off",
                      Rect(x0 + (k - 1) * (size + gap), y, size, size), f"tracker_{k}") for k in range(1, n + 1)]

    # ---- filling ----------------------------------------------------------------------------------------------

    def _hero_number(self, content: _Content) -> Fact | None:
        """The first number fact the single item cites: shown as the hero value of a single-hero slide."""
        if len(content.items) != 1:
            return None
        return next((self.facts[f] for f in content.items[0].fact_ids if self.facts[f].number_value), None)

    def _fills(self, layout: Any, content: _Content, spec: dict) -> dict[str, list[str]]:
        names = {ph.name for ph in layout.placeholders}
        title = content.title
        if content.part is not None:
            title = self.labels["part_title"].format(title=title, k=content.part[0], n=content.part[1])
        fills: dict[str, list[str]] = {"title": [title]}
        slots = spec.get("slots", {})
        stems = slots.get("stack", []) + slots.get("left", []) + slots.get("right", [])

        if content.slide_type in (SlideType.TITLE, SlideType.SECTION_DIVIDER):
            if content.items and "subtitle" in names:
                fills["subtitle"] = [item.text for item in content.items]
        elif content.lines:
            fills["body"] = list(content.lines)
        elif "item" in stems:
            takeaways, decisions = self._split_decisions(content)
            slot_items = takeaways if "decisions" in names else list(content.items)
            for n, item in enumerate(slot_items, start=1):
                lead = _LEAD.match(item.text)
                if "item_title" in stems and f"item_{n}_title" in names:
                    fills[f"item_{n}_title"] = [lead["lead"] if lead else self.labels["item_ordinal"].format(n=n)]
                    fills[f"item_{n}"] = [lead["rest"] if lead else item.text]
                else:
                    fills[f"item_{n}"] = [item.text]
                if "marker" in stems:
                    fills[f"marker_{n}"] = [self.labels["marker"].format(n=n)]
            if decisions and "decisions" in names:
                fills["decisions_title"] = [self.labels["decisions_needed"]]
                fills["decisions"] = [item.text for item in decisions]
        elif content.items and "body" in names:
            fills["body"] = [item.text for item in content.items]

        if content.section_number is not None:
            fills["section_number"] = [self.labels["section_number"].format(n=content.section_number)]
        for n, fact_id in enumerate(content.kpi_fact_ids, start=1):
            value = self.facts[fact_id].number_value
            fills[f"kpi_value_{n}"] = [value.value_as_written]
            fills[f"kpi_label_{n}"] = [value.label]
        panel = spec.get("hero_panel")
        hero = self._hero_number(content) if panel else None
        if hero is not None:
            fills["kpi_value_1"] = [hero.number_value.value_as_written]
            fills["kpi_label_1"] = [hero.number_value.label]
        elif "quote" in names and content.items and (panel or "quote" in spec.get("fixed", {})):
            evidence = self.facts[content.items[0].fact_ids[0]].evidence[0]
            fills["quote"] = [f"“{evidence.quote}”"]
            if "attribution" in names:
                fills["attribution"] = [self.labels["page"].format(n=page_number(evidence.block_id))]
        if "source" in names:
            fills["source"] = [self._pages_text(content) or self.labels["no_pages"]]
        return fills

    def _kpi_width(self, value: str, unit: str, style: TextStyle, unit_pt: float | None) -> float:
        """Width of a KPI value with its unit run (units are body size, trick 17)."""
        small = replace(style, size_pt=unit_pt or style.size_pt / 2)
        return self.fitter.width(value, style) + (self.fitter.width(unit, small) if unit else 0.0)

    def _line_boxes(self, paragraphs: list[str], style: TextStyle, rect: Rect, top: float) -> tuple[Rect, ...]:
        """Ink of a text box: one rectangle per wrapped line, at its measured width (paragraph spacing ignored)."""
        inner = (rect.w - style.insets[0] - style.insets[2] - style.indent_emu) * WIDTH_SAFETY
        line_h = self.fitter.line_height(style)
        boxes, y = [], top + style.insets[1]
        x0 = rect.x + style.insets[0] + style.indent_emu
        for text in paragraphs:
            for line in self.fitter.wrap(text, inner, style):
                w = min(inner / WIDTH_SAFETY, self.fitter.width(line, style) / WIDTH_SAFETY)
                x = rect.x + (rect.w - w) / 2 if style.align == "ctr" else x0
                if w > 0:
                    boxes.append(Rect(x, y, w, line_h))
                y += line_h
        return tuple(boxes)

    def _unit(self, fact_id: str) -> str:
        unit = self.facts[fact_id].number_value.unit
        return "" if unit is None else unit if unit.startswith("%") else f" {unit}"

    def _pages_text(self, content: _Content) -> str:
        pages = {page_number(ev.block_id) for fid in content.fact_ids for ev in self.facts[fid].evidence}
        entries = [(n, self.labels["page"].format(n=n)) for n in sorted(pages)]
        if content.chart is not None:
            table = self.tables[content.chart.table_id]
            entries.append((table.page_index + 1,
                            self.labels["page"].format(n=table.page_index + 1) + self.labels["table_suffix"]))
        if not entries:
            return ""
        texts = [text for _, text in sorted(entries)]
        first = texts[0]
        rest = [t.removeprefix(self.labels["page"].format(n="")) for t in texts[1:]]
        return self.labels["pages_joiner"].join([first, *rest])

    def _appendix(self, content: _Content, cited: list[str]) -> _Content:
        lines, ids = [], []
        for fact_id in cited:
            fact = self.facts[fact_id]
            if fact.number_value is None:
                continue
            value = fact.number_value
            for ev in fact.evidence:
                lines.append(self.labels["appendix_line"].format(
                    value=value.value_as_written + self._unit(fact_id), label=value.label,
                    page=self.labels["page"].format(n=page_number(ev.block_id)), quote=ev.quote))
                ids.append(fact_id)
        if not lines:
            lines = [self.labels["appendix_empty"]]
        return replace(content, lines=tuple(lines), line_fact_ids=tuple(ids))

    def _sources_lines(self, plan: DeckPlan, planned: list[list[_Plan]], cited: list[str]) -> list[str]:
        names = self.labels["slide_type_names"]
        lines = [item.text for slide in plan.slides if slide.slide_type is SlideType.SOURCES for item in slide.items]
        number = 0
        for slide, parts in zip(plan.slides, planned):
            for part in parts:
                number += 1
                if slide.slide_type in (SlideType.SOURCES, SlideType.TITLE, SlideType.SECTION_DIVIDER):
                    continue
                pages = self._pages_text(part.content) or self.labels["no_pages"]
                lines.append(self.labels["sources_line"].format(n=number, name=names[slide.slide_type.value],
                                                                pages=pages))
        verified = sum(self.facts[fid].status is FactStatus.VERIFIED for fid in cited)
        all_pages = sorted({page_number(ev.block_id) for fid in cited for ev in self.facts[fid].evidence}
                           | {self.tables[s.chart.table_id].page_index + 1 for s in plan.slides if s.chart})
        lines += [
            self.labels["facts_verified"].format(verified=verified, total=len(cited)),
            self.labels["pages_cited"].format(pages=self.labels["pages_joiner"].join(map(str, all_pages))),
            self.labels["document"].format(doc=f"{plan.doc_id[:19]}…"),
            self.labels["method"],
        ]
        return lines

    # ---- charts -------------------------------------------------------------------------------------------------

    def _decide_chart(self, content: _Content) -> ChartDecision:
        spec = content.chart
        cited = [self.facts[fid] for fid in content.fact_ids]
        try:
            return decide(self.tables[spec.table_id], spec, cited, series_fallback=self.labels["series_fallback"])
        except ChartDataError as exc:
            raise DeckRenderError(str(exc)) from exc

    def _chart(self, placeholder: Any, plan: _Plan) -> None:
        if placeholder.placeholder_format.type != PP_PLACEHOLDER.CHART:
            raise DeckRenderError("CHART_PLACEHOLDER_WRONG_KIND")
        decision, spec = plan.chart, plan.content.chart
        if decision.kind != spec.chart_type.value:
            warning = f"CHART_TYPE_BY_RULE:{spec.table_id}:{spec.chart_type.value}->{decision.kind}"
            if warning not in self.warnings:
                self.warnings.append(warning)
        rules = load_rules()
        chart_data = CategoryChartData(number_format=decision.number_format)
        stacked = decision.kind == "stacked_bar"
        if stacked:  # one bar; each part is a series, labelled with its own name and value
            chart_data.categories = [""]
            for name, value in zip(decision.categories, decision.series[0].values):
                chart_data.add_series(name, [float(value)])
        else:
            chart_data.categories = list(decision.categories)
            for series in decision.series:
                chart_data.add_series(series.name, [float(v) for v in series.values])
        frame = placeholder.insert_chart(_CHART_TYPES[decision.kind], chart_data)
        frame.name = "chart"
        chart = frame.chart
        chart.font.size = Pt(plan.body_pt or rules["font_pt"])  # chart text is body size (trick 17)
        chart.has_title = False
        chart.has_legend = len(decision.series) >= rules["labels"]["legend_min_series"] and not stacked
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.vary_by_categories = False
        if decision.kind != "line":
            plot.gap_width = 60
            if stacked:
                plot.overlap = 100
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = decision.number_format
        labels.number_format_is_linked = False
        labels.position = (XL_LABEL_POSITION.ABOVE if decision.kind == "line" else XL_LABEL_POSITION.CENTER
                           if stacked else XL_LABEL_POSITION.OUTSIDE_END)
        if stacked:
            labels.show_series_name = True
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.has_minor_gridlines = False
        chart.value_axis.visible = False
        chart.category_axis.has_major_gridlines = False
        if decision.kind in ("bar", "stacked_bar"):
            chart.category_axis.reverse_order = True  # first row on top, as read (sorted: largest on top)
        if stacked:
            chart.category_axis.visible = False
        self._chart_colors(chart, decision, rules, stacked)
        self._callout(chart, decision, rules, stacked, plan.body_pt or rules["callout_pt"])

    def _chart_colors(self, chart: Any, decision: ChartDecision, rules: dict, stacked: bool) -> None:
        """Trick 11: the key point in the accent, the rest muted; several series take a colorblind-safe order that
        starts from the accent, with the key series first."""
        muted = decor_mod.tint("tx1", rules["focus"]["muted_alpha"])
        palette = [self.focus_xml if slot == "accent1" else muted if slot == "muted" else decor_mod.tint(slot, 1.0)
                   for slot in rules["palette"]]
        plot = chart.plots[0]
        key_series, key_point = decision.key
        if stacked:
            for index, series in enumerate(plot.series):
                alpha = rules["focus"]["muted_alpha"] + 0.15 * (index % 2)
                _set_fill(series.format, self.focus_xml if index == key_point else decor_mod.tint("tx1", alpha))
                _set_line(series.format, decor_mod.tint("bg1", 1.0), 1.5)
            return
        order = [key_series] + [i for i in range(len(decision.series)) if i != key_series]
        colors = {s: palette[min(rank, len(palette) - 1)] for rank, s in enumerate(order)}
        for index, series in enumerate(plot.series):
            if decision.kind == "line":
                _set_line(series.format, colors[index], 3)
                _set_fill(series.marker.format, colors[index])
                continue
            _set_fill(series.format, colors[index])
            if len(decision.series) == 1:
                for p in range(len(decision.categories)):
                    _set_fill(series.points[p].format, self.focus_xml if p == key_point else muted)

    def _callout(self, chart: Any, decision: ChartDecision, rules: dict, stacked: bool, size_pt: float) -> None:
        """Trick 13: one annotation on the key point, copied from the cited fact (its value, plus its label when the
        label is short enough to sit on a chart), otherwise naming the maximum's value as the table writes it."""
        s_index, p_index = decision.key
        if decision.key_fact_id is not None:
            fact = self.facts[decision.key_fact_id]
            value = fact.number_value.value_as_written + self._unit(fact.fact_id)
            label = fact.number_value.label
            text = (self.labels["callout_cited"].format(label=label, value=value)
                    if len(label) <= rules["callout_label_max"] else value)
        else:
            text = self.labels["callout_max"].format(value=decision.series[s_index].texts[p_index])
        series = chart.plots[0].series[p_index if stacked else s_index]
        label = series.points[0 if stacked else p_index].data_label
        label.has_text_frame = True
        label.text_frame.text = text
        run = label.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(size_pt)
        _set_run_color(run, self.accent_xml)
        label.position = (XL_LABEL_POSITION.ABOVE if decision.kind == "line" else XL_LABEL_POSITION.CENTER
                          if stacked else XL_LABEL_POSITION.OUTSIDE_END)

    # ---- emitting -----------------------------------------------------------------------------------------------

    def _emit(self, plan: _Plan) -> None:
        slide = self.prs.slides.add_slide(plan.layout)
        number = len(self.prs.slides)
        self.layouts_used.append(plan.layout.name)
        self.overflows += [replace(o, slide_number=number) for o in plan.overflows]
        if plan.score is not None and plan.score.hard:
            self.warnings += [f"HARD_VIOLATION:slide {number}:{h}" for h in plan.score.hard]
        self.scores.append({"slide": number, "plan_index": plan.content.plan_index, "part": plan.content.part,
                            **plan.record_, "chosen": plan.record()})
        layout_names = {ph.placeholder_format.idx: ph.name for ph in plan.layout.placeholders}
        for shape in list(slide.placeholders):
            name = layout_names.get(shape.placeholder_format.idx)
            shape.name = name or shape.name
            rect = plan.rects.get(name)
            if rect is not None and (name in plan.fills or name == "chart"):
                shape.left, shape.top = int(rect.x), int(rect.y)
                shape.width, shape.height = int(rect.w), int(rect.h)
            if name == "chart" and plan.chart is not None:
                self._chart(shape, plan)
            elif name == "title":
                self._fill_title(shape, plan)
            elif name in plan.fills:
                self._fill_text(shape, plan.fills[name], plan.sizes.get(name))
                if name in plan.anchors and name != "source":
                    shape.text_frame.vertical_anchor = (MSO_ANCHOR.MIDDLE if plan.anchors[name] == "center"
                                                        else MSO_ANCHOR.TOP)
                if name in plan.units:
                    value = plan.sizes.get(name) or self._style(plan.layout, self._layout_ph(plan.layout, name)).size_pt
                    run = shape.text_frame.paragraphs[0].add_run()
                    run.text = plan.units[name]
                    run.font.size = int((plan.body_pt or value / 2) * EMU_PER_PT)
                if name == "source":
                    self._fit_chip(shape, plan)
                if name == "section_number":
                    decor_mod.style_divider_number(shape, self.divider_style, self.accent_xml)
            else:
                shape._element.getparent().remove(shape._element)
        decor_mod.draw(slide, plan.decor, self.focus_xml, self.card_style)
        self._notes(slide, plan.content)

    def _fill_title(self, shape: Any, plan: _Plan) -> None:
        """Trick 18 and sentence case: balanced line breaks by measured width; theme caps only on short titles."""
        text = plan.fills["title"][0]
        ph = self._layout_ph(plan.layout, "title")
        base = self._style(plan.layout, ph)
        style = self._title_style(plan.layout, ph, plan.sizes.get("title") or base.size_pt, text)
        inner = (ph.width - style.insets[0] - style.insets[2] - style.indent_emu) * WIDTH_SAFETY
        lines = balanced_lines(text, inner, lambda t: self.fitter.width(t, style), self.fitter.wrap(text, inner, style))
        paragraph = shape.text_frame.paragraphs[0]
        offset = 0
        for index, line in enumerate(lines):
            if index:
                paragraph.add_line_break()
                offset += 1  # the space the break replaces
            pieces, cursor = [], 0
            for start, end in plan.title_accents:  # trick 16: split the line into plain and accent runs
                a, b = start - offset, end - offset
                if 0 <= a and b <= len(line):
                    pieces += [(line[cursor:a], False), (line[a:b], True)]
                    cursor = b
            pieces.append((line[cursor:], False))
            for text_piece, accent in pieces:
                if not text_piece:
                    continue
                run = paragraph.add_run()
                run.text = text_piece
                if "title" in plan.sizes:
                    run.font.size = int(plan.sizes["title"] * EMU_PER_PT)
                if base.caps and not style.caps:
                    run.font._rPr.set("cap", "none")
                if accent:
                    _set_run_color(run, self.accent_xml)
            offset += len(line)

    def _fill_text(self, shape: Any, paragraphs: list[str], size: float | None) -> None:
        frame = shape.text_frame
        for index, text in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = text
            if size:
                run.font.size = int(size * EMU_PER_PT)

    @staticmethod
    def _layout_ph(layout: Any, name: str) -> Any:
        return next(ph for ph in layout.placeholders if ph.name == name)

    def _fit_chip(self, shape: Any, plan: _Plan) -> None:
        ph = self._layout_ph(plan.layout, "source")
        width = self.fitter.text_width(plan.fills["source"][0], self._style(plan.layout, ph))
        left, top, height = shape.left, shape.top, shape.height
        shape.left, shape.top, shape.width, shape.height = left, top, max(width, height * 2), height

    def _notes(self, slide: Any, content: _Content) -> None:
        lines = [content.notes, ""] if content.notes else []
        table = self.tables.get(content.chart.table_id) if content.chart else None
        if content.fact_ids or table is not None:
            lines.append(self.labels["notes_heading"])
            for fid in content.fact_ids:
                fact = self.facts[fid]
                for ev in fact.evidence:
                    page = self.labels["page"].format(n=page_number(ev.block_id))
                    lines.append(self.labels["notes_line"].format(fact_id=fid, page=page, status=fact.status.value,
                                                                  quote=ev.quote))
            if table is not None:
                page = self.labels["page"].format(n=table.page_index + 1)
                lines.append(self.labels["notes_table"].format(table_id=table.table_id, page=page))
        if lines:
            slide.notes_slide.notes_text_frame.text = "\n".join(lines)


@dataclass(frozen=True)
class _Part:
    """One entry of a slot stack: a placeholder (with its type-scale role) or a native block."""
    name: str
    role: str | None
    kind: str  # text, kpi, hero, quote, marker, block
    slot: int = 0


class _Measure:
    """Measures and places one candidate's text; records sizes, boxes and anchors on its plan."""

    def __init__(self, builder: _DeckBuilder, plan: _Plan, step: dict, scored: bool) -> None:
        self.b = builder
        self.plan = plan
        self.layout = plan.layout
        self.by_name = {ph.name: ph for ph in plan.layout.placeholders}
        self.step = step
        self.scored = scored
        self.snap_values = False

    def size_for(self, role: str | None) -> float | None:
        """The candidate's size for a role, snapped to the template's three content sizes (trick 17)."""
        if not self.scored or role is None:
            return None
        factor = self.step["factor"]
        sizes = self.b.type_sizes
        if role == "hero":
            return float(max(self.b.cfg["hero_min_pt"], round(sizes["hero"] * factor)))
        if role in VALUE_ROLES:
            return float(self.b.title_pt) if self.snap_values else float(round(sizes["display"] * factor))
        return float(round(sizes["body"] * factor))

    def style(self, name: str, role: str | None) -> TextStyle:
        style = self.b._style(self.layout, self.by_name[name])
        size = self.size_for(role)
        return replace(style, size_pt=size) if size else style

    def fit_values(self, columns: list[tuple[list[_Part], Rect]]) -> None:
        """KPI values keep the display size when every one (with its unit) fits its column on one line; otherwise
        all of them snap to the title size, so the slide still uses three sizes."""
        values = [(p, column) for parts, column in columns for p in parts if p.role == "kpi_value"]
        self.snap_values = False
        if not all(self._value_fits(p, column) for p, column in values):
            self.snap_values = True

    def _value_fits(self, part: _Part, column: Rect) -> bool:
        style = self.style(part.name, part.role)
        inner = (column.w - style.insets[0] - style.insets[2]) * WIDTH_SAFETY
        unit = self.plan.units.get(part.name, "")
        return self.b._kpi_width(self.plan.fills[part.name][0], unit, style, self.plan.body_pt) <= inner

    def part_height(self, part: _Part, width: float) -> float:
        if part.kind == "block":
            return self.b._block_height(part, width)
        if part.kind == "marker":
            return min(width, self.b.geo["marker_pt"] * EMU_PER_PT)
        return self.b._extent(self.plan.fills[part.name], self.style(part.name, part.role), width)[0]

    def height_of(self, parts: list[_Part], column: Rect) -> float:
        if not parts:
            return 0.0
        gap = self.b.geo["stack_gap_pt"] * EMU_PER_PT
        return sum(self.part_height(p, column.w) for p in parts) + gap * (len(parts) - 1)

    def offset(self, parts: list[_Part], column: Rect, anchor: str) -> float:
        return 0.0 if anchor == "top" else max(0.0, (column.h - self.height_of(parts, column)) / 2)

    def stack(self, parts: list[_Part], column: Rect, offset: float, backdrop: str | None) -> list[Element]:
        """Place `parts` top-down in `column`, `offset` below its top, each at its measured height."""
        gap = self.b.geo["stack_gap_pt"] * EMU_PER_PT
        slack = self.b.geo["slack_pt"] * EMU_PER_PT
        elements: list[Element] = []
        y = column.y + offset
        text_parts = [p for p in parts if p.kind not in ("block", "marker")]
        insets = self.style(text_parts[0].name, text_parts[0].role).insets if text_parts else (0, 0, 0, 0)
        for part in parts:
            h = self.part_height(part, column.w)
            beyond = y + h > column.bottom + EMU_PER_PT
            if part.kind == "block":
                rect = Rect(column.x + insets[0], y, column.w - insets[0] - insets[2], h)
                if beyond:
                    elements.append(Element(f"{part.name}_{part.slot}", "text", rect, overflow=True))
                elements += self.b._place_block(self, part, rect, backdrop)
            elif part.kind == "marker":
                elements.append(self.place(part.name, Rect(column.x, y, h, h), part.role, "accent", "center",
                                           "marker"))
            else:
                rect = Rect(column.x, y, column.w, min(h + slack, max(h, column.bottom - y)))
                element = self.place(part.name, rect, part.role, backdrop, "top", part.kind)
                if beyond and not element.overflow:
                    element = replace(element, overflow=True)
                    self.plan.overflows.append(Overflow(0, part.name, int(y + h - column.y), int(column.h)))
                elements.append(element)
            y += h + gap
        return elements

    def place(self, name: str, rect: Rect, role: str | None, backdrop: str | None, anchor: str = "top",
              kind: str = "text") -> Element:
        style = self.style(name, role)
        if self.size_for(role):
            self.plan.sizes[name] = style.size_pt
        paragraphs = self.plan.fills[name]
        height, width = self.b._extent(paragraphs, style, rect.w)
        overflow = height > rect.h + EMU_PER_PT
        if overflow:
            self.plan.overflows.append(Overflow(0, name, int(height), int(rect.h)))
        unit = self.plan.units.get(name)
        if unit is not None:  # value and unit stay on one line
            value_w = self.b._kpi_width(paragraphs[0], unit, style, self.plan.body_pt)
            inner = (rect.w - style.insets[0] - style.insets[2]) * WIDTH_SAFETY
            width = min(rect.w, value_w / WIDTH_SAFETY + style.insets[0] + style.insets[2])
            if value_w > inner:
                overflow = True
                self.plan.overflows.append(Overflow(0, name, int(value_w), int(inner)))
        top = rect.y + max(0.0, (rect.h - height) / 2) if anchor == "center" else rect.y
        if unit is not None:
            ink = (Rect(rect.x + style.insets[0], top + style.insets[1], width - style.insets[0] - style.insets[2],
                        self.b.fitter.line_height(style)),)
        else:
            ink = self.b._line_boxes(paragraphs, style, rect, top)
        x = rect.x + (rect.w - width) / 2 if style.align == "ctr" else rect.x
        self.plan.rects[name] = rect
        self.plan.anchors[name] = anchor
        if backdrop == "accent":
            ratio = contrast(self.b._fg(style), self.b.colors.resolve(self.b.accent_xml)[0])
        else:
            ratio = self.b._contrast(style, self.layout, backdrop)
        if name == "section_number" and self.b.divider_style.get("type") == "outline_number":
            foreground = hex_rgb(self.b.divider_style["outline_color"])
            ratio = min(contrast(foreground, background) for background in self.b.colors.background(self.layout))
        return Element(name, kind, Rect(x, top, width, min(height, rect.h)), ink=0.5 if kind == "text" else 0.85,
                       role=role, size_pt=style.size_pt, bold=style.bold, overflow=overflow, contrast=ratio,
                       ink_rects=ink)


def _set_fill(fmt: Any, color_xml: str) -> None:
    fmt.fill.solid()
    solid = fmt.fill._xPr.find(f"{{{_A}}}solidFill")
    for child in list(solid):
        solid.remove(child)
    solid.append(etree.fromstring(color_xml))


def _set_line(fmt: Any, color_xml: str, width_pt: float) -> None:
    line = fmt.line
    line.width = Pt(width_pt)
    line.fill.solid()
    solid = line._ln.find(f"{{{_A}}}solidFill")
    for child in list(solid):
        solid.remove(child)
    solid.append(etree.fromstring(color_xml))


def _set_run_color(run: Any, color_xml: str) -> None:
    rpr = run.font._rPr
    for child in rpr.findall(f"{{{_A}}}solidFill"):
        rpr.remove(child)
    solid = etree.Element(f"{{{_A}}}solidFill")
    solid.append(etree.fromstring(color_xml))
    rpr.insert(0, solid)
