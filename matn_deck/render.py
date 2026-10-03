"""Renderer v0: DeckPlan + facts + tables -> editable .pptx, filling only the slots in design/layouts.json.

Every color, font and size comes from design/tokens.json. Text is measured with the real font files and overflow is
reported, never auto-shrunk. KPI values and chart data are copied from facts and tables by code.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
import re
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from matn_contracts.domain import ChartSpec, DeckPlan, Fact, Slide, Table, validate_deck_plan_against_facts
from matn_contracts.enums import ChartType, FactStatus, SlideType

from .design import Box, Design, load_design
from .measure import EMU_PER_PT, TextMeasurer

_PAGE_IN_ID = re.compile(r":p(\d{4,}):")
_NUMBER = re.compile(r"(?P<sign>[-−])?(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?:\.(?P<frac>\d+))?(?P<pct>%)?")
_ALIGN = {"start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_CHART_TYPES = {
    ChartType.BAR: XL_CHART_TYPE.BAR_CLUSTERED,
    ChartType.COLUMN: XL_CHART_TYPE.COLUMN_CLUSTERED,
    ChartType.LINE: XL_CHART_TYPE.LINE_MARKERS,
}


class DeckRenderError(ValueError):
    pass


@dataclass(frozen=True)
class Overflow:
    slide_number: int
    slot: str
    dimension: str  # "height" or "width"
    needed_emu: int
    available_emu: int


@dataclass(frozen=True)
class RenderReport:
    overflows: tuple[Overflow, ...]
    warnings: tuple[str, ...]


class DeckOverflowError(DeckRenderError):
    def __init__(self, report: RenderReport) -> None:
        slots = ", ".join(f"slide {o.slide_number} {o.slot}" for o in report.overflows)
        super().__init__(f"TEXT_OVERFLOW: {slots}")
        self.report = report


@dataclass(frozen=True)
class _Run:
    text: str
    role: str
    size_pt: float | None = None


@dataclass(frozen=True)
class _Para:
    runs: tuple[_Run, ...]
    bullet: str | None = None


@dataclass(frozen=True)
class _Provenance:
    fact_ids: tuple[str, ...]
    table: Table | None


def page_number(block_or_table_id: str) -> int:
    """1-based physical PDF page encoded in a block_id or table_id."""
    match = _PAGE_IN_ID.search(block_or_table_id)
    if match is None:
        raise DeckRenderError(f"ID_HAS_NO_PAGE:{block_or_table_id}")
    return int(match.group(1)) + 1


def build_presentation(
    plan: DeckPlan,
    facts: tuple[Fact, ...],
    tables: tuple[Table, ...],
    *,
    strict: bool = False,
    design: Design | None = None,
) -> tuple[Any, RenderReport]:
    builder = _DeckBuilder(design or load_design(), facts, tables)
    prs = builder.build(plan)
    report = RenderReport(overflows=tuple(builder.overflows), warnings=tuple(builder.warnings))
    if strict and report.overflows:
        raise DeckOverflowError(report)
    return prs, report


def render(
    plan: DeckPlan,
    facts: tuple[Fact, ...],
    tables: tuple[Table, ...],
    out_path: Path,
    *,
    strict: bool = False,
    design: Design | None = None,
) -> RenderReport:
    prs, report = build_presentation(plan, facts, tables, strict=strict, design=design)
    prs.save(str(out_path))
    return report


class _DeckBuilder:
    def __init__(self, design: Design, facts: tuple[Fact, ...], tables: tuple[Table, ...]) -> None:
        self.d = design
        self.m = TextMeasurer(design)
        self.facts = {fact.fact_id: fact for fact in facts}
        self.tables = {table.table_id: table for table in tables}
        self.overflows: list[Overflow] = []
        self.warnings: list[str] = []
        self.number = 0

    # ---- deck ---------------------------------------------------------------------------------------------------

    def build(self, plan: DeckPlan) -> Any:
        validate_deck_plan_against_facts(plan, tuple(self.facts.values()))
        provenance = [self._provenance(slide) for slide in plan.slides]
        for prov in provenance:
            for fact_id in prov.fact_ids:
                status = self.facts[fact_id].status
                if status is FactStatus.REJECTED:
                    raise DeckRenderError(f"REJECTED_FACT_CITED:{fact_id}")
                if status is not FactStatus.VERIFIED and f"UNVERIFIED_FACT:{fact_id}" not in self.warnings:
                    self.warnings.append(f"UNVERIFIED_FACT:{fact_id}")

        prs = Presentation()
        prs.slide_width = self.d.tokens["slide"]["width_emu"]
        prs.slide_height = self.d.tokens["slide"]["height_emu"]
        prs.core_properties.title = plan.slides[0].title
        blank = prs.slide_layouts[6]
        handlers = {
            SlideType.TITLE: self._title,
            SlideType.EXECUTIVE_SUMMARY: self._executive_summary,
            SlideType.KPI: self._kpi,
            SlideType.FINDING: self._finding,
            SlideType.CHART: self._chart_slide,
            SlideType.IMPLICATIONS: self._implications,
            SlideType.LIMITATIONS: self._limitations,
            SlideType.SOURCES: lambda s, p, a, prov: self._sources(s, p, a, plan, provenance),
        }
        for index, (slide_plan, prov) in enumerate(zip(plan.slides, provenance), start=1):
            self.number = index
            arch = self.d.archetype(slide_plan.slide_type.value)
            slide = prs.slides.add_slide(blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = self._rgb(arch.get("background", self.d.layouts["background"]))
            if arch.get("frame"):
                self._frame(slide, slide_plan, arch, prov)
            handlers[slide_plan.slide_type](slide, slide_plan, arch, prov)
            self._notes(slide, slide_plan, prov)
        return prs

    def _provenance(self, slide: Slide) -> _Provenance:
        fact_ids = [fid for item in slide.items for fid in item.fact_ids] + list(slide.kpi_fact_ids)
        table = None
        if slide.chart is not None:
            table = self.tables.get(slide.chart.table_id)
            if table is None:
                raise DeckRenderError(f"UNKNOWN_TABLE_ID:{slide.chart.table_id}")
        return _Provenance(tuple(dict.fromkeys(fact_ids)), table)

    def _pages(self, prov: _Provenance) -> list[str]:
        pages = {page_number(ev.block_id) for fid in prov.fact_ids for ev in self.facts[fid].evidence}
        entries = [(n, 0, self.d.label("page", n=n)) for n in pages]
        if prov.table is not None:
            entries.append((prov.table.page_index + 1, 1, self.d.label("table_page", n=prov.table.page_index + 1)))
        return [text for _, _, text in sorted(entries)]

    # ---- archetypes ---------------------------------------------------------------------------------------------

    def _frame(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        frame = self.d.layouts["frame"]
        self._rect(slide, "accent_bar", frame["accent_bar"])
        self._text(slide, "title", frame["title"], [self._para(plan.title, frame["title"]["role"])])
        if arch.get("source_footer", True):
            pages = self._pages(prov)
            joined = self.d.layouts["labels"]["pages_joiner"].join(pages) or self.d.label("no_pages")
            self._text(slide, "footer_source", frame["footer_source"],
                       [self._para(self.d.label("source", pages=joined), "footer")])
        self._text(slide, "footer_page", frame["footer_page"], [self._para(str(self.number), "footer")])

    def _title(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        slots = arch["slots"]
        self._rect(slide, "accent_bar", slots["accent_bar"])
        self._text(slide, "title", slots["title"], [self._para(plan.title, slots["title"]["role"])])
        if plan.items:
            role = slots["subtitle"]["role"]
            self._text(slide, "subtitle", slots["subtitle"], [self._para(item.text, role) for item in plan.items])

    def _executive_summary(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        self._numbered(slide, plan, arch["repeat"]["cells"], arch["repeat"]["slots"], "rule")

    def _implications(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        self._numbered(slide, plan, arch["repeat"]["cells"], arch["repeat"]["slots"], "card")

    def _numbered(self, slide: Any, plan: Slide, cells: list, slots: dict, rect: str) -> None:
        for n, (item, cell) in enumerate(zip(plan.items, cells), start=1):
            self._rect(slide, f"item_{n}_{rect}", slots[rect], cell)
            self._text(slide, f"item_{n}_number", slots["number"], [self._para(str(n), slots["number"]["role"])],
                       cell)
            self._text(slide, f"item_{n}_text", slots["text"], [self._para(item.text, slots["text"]["role"])], cell)

    def _kpi(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        slots = arch["repeat"]["slots"]
        cells = arch["repeat"]["cells_by_count"][str(len(plan.kpi_fact_ids))]
        values = [self.facts[fid].number_value for fid in plan.kpi_fact_ids]
        unit_role = self.d.role("kpi_unit")
        value_role = self.d.role(slots["value"]["role"])

        def unit_text(unit: str | None) -> str:
            return "" if unit is None else unit if unit.startswith("%") else f" {unit}"

        def value_width(value: Any, size: float) -> float:
            return (self.m.width(value.value_as_written, size, value_role["bold"], value_role["italic"])
                    + self.m.width(unit_text(value.unit), unit_role["size_pt"], unit_role["bold"], unit_role["italic"]))

        box = self.d.resolve_box(slots["value"], cells[0])
        available = box.w * self.m.width_safety
        steps = self.d.tokens["kpi_value_steps_pt"]
        size = next((s for s in steps if all(value_width(v, s) <= available for v in values)
                     and self.m.line_height(s, value_role["line_spacing"]) <= box.h), steps[-1])
        for n, (fact_id, value, cell) in enumerate(zip(plan.kpi_fact_ids, values, cells), start=1):
            self._rect(slide, f"kpi_{n}_tile", slots["tile"], cell)
            needed = value_width(value, size)
            if needed > available:
                self.overflows.append(Overflow(self.number, f"kpi_{n}_value", "width", int(needed), int(available)))
            runs = [_Run(value.value_as_written, "kpi_value", size)]
            if value.unit is not None:
                runs.append(_Run(unit_text(value.unit), "kpi_unit"))
            self._text(slide, f"kpi_{n}_value", slots["value"], [_Para(tuple(runs))], cell, measure=False)
            self._text(slide, f"kpi_{n}_label", slots["label"], [self._para(value.label, slots["label"]["role"])],
                       cell)
            page = self.d.label("page", n=page_number(self.facts[fact_id].evidence[0].block_id))
            self._text(slide, f"kpi_{n}_source", slots["source"], [self._para(page, slots["source"]["role"])], cell)
        if plan.items:
            self._list(slide, "items", arch["slots"]["items"], [item.text for item in plan.items])

    def _finding(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        slots = arch["slots"]
        if not plan.items:
            return
        self._list(slide, "items", slots["items"], [item.text for item in plan.items])
        evidence = self.facts[plan.items[0].fact_ids[0]].evidence[0]
        self._rect(slide, "panel", slots["panel"])
        heading = self.d.label("evidence_heading", page=self.d.label("page", n=page_number(evidence.block_id)))
        self._text(slide, "panel_heading", slots["panel_heading"], [self._para(heading, "panel_heading")])
        self._text(slide, "quote", slots["quote"], [self._para(f"“{evidence.quote}”", "quote")])

    def _chart_slide(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        slots = arch["slots"]
        self._chart(slide, self.d.resolve_box(slots["chart"]), plan.chart, prov.table)
        if plan.items:
            self._list(slide, "items", slots["items"], [item.text for item in plan.items])

    def _limitations(self, slide: Any, plan: Slide, arch: dict, prov: _Provenance) -> None:
        self._list(slide, "items", arch["slots"]["items"], [item.text for item in plan.items])

    def _sources(self, slide: Any, plan: Slide, arch: dict, deck: DeckPlan, provenance: list[_Provenance]) -> None:
        slots = arch["slots"]
        role = slots["list"]["role"]
        names = self.d.layouts["labels"]["archetype_names"]
        paras = [self._para(item.text, role) for item in plan.items]
        for n, (other, prov) in enumerate(zip(deck.slides, provenance), start=1):
            if other is plan:
                continue
            pages = self.d.layouts["labels"]["pages_joiner"].join(self._pages(prov)) or self.d.label("no_pages")
            paras.append(self._para(self.d.label("sources_line", n=n, name=names[other.slide_type.value],
                                                 pages=pages), role))
        self._text(slide, "list", slots["list"], paras)

        cited = list(dict.fromkeys(fid for prov in provenance for fid in prov.fact_ids))
        verified = sum(self.facts[fid].status is FactStatus.VERIFIED for fid in cited)
        all_pages = sorted({page_number(ev.block_id) for fid in cited for ev in self.facts[fid].evidence}
                           | {prov.table.page_index + 1 for prov in provenance if prov.table is not None})
        joiner = self.d.layouts["labels"]["pages_joiner"]
        body_role = slots["panel_body"]["role"]
        self._rect(slide, "panel", slots["panel"])
        self._text(slide, "panel_heading", slots["panel_heading"],
                   [self._para(self.d.label("verification_heading"), "panel_heading")])
        self._text(slide, "panel_body", slots["panel_body"], [
            self._para(self.d.label("facts_verified", verified=verified, total=len(cited)), body_role),
            self._para(self.d.label("pages_cited", pages=joiner.join(str(p) for p in all_pages)), body_role),
            self._para(self.d.label("document", doc=f"{deck.doc_id[:19]}…"), body_role),
            self._para(self.d.label("method"), body_role),
        ])

    def _notes(self, slide: Any, plan: Slide, prov: _Provenance) -> None:
        lines = [plan.notes, ""] if plan.notes else []
        if prov.fact_ids or prov.table is not None:
            lines.append(self.d.label("notes_heading"))
            for fid in prov.fact_ids:
                fact = self.facts[fid]
                for ev in fact.evidence:
                    page = self.d.label("page", n=page_number(ev.block_id))
                    lines.append(self.d.label("notes_line", fact_id=fid, page=page, status=fact.status.value,
                                              quote=ev.quote))
            if prov.table is not None:
                page = self.d.label("page", n=prov.table.page_index + 1)
                lines.append(self.d.label("notes_table", table_id=prov.table.table_id, page=page))
        if lines:
            slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    # ---- primitives ---------------------------------------------------------------------------------------------

    def _rgb(self, token: str) -> RGBColor:
        return RGBColor.from_string(self.d.color(token))

    def _para(self, text: str, role: str, bullet: str | None = None) -> _Para:
        return _Para((_Run(text, role),), bullet)

    def _list(self, slide: Any, name: str, slot: dict, texts: list[str]) -> None:
        bullet = self.d.tokens["bullet"]["caveat_char" if slot.get("list") == "caveat" else "char"]
        self._text(slide, name, slot, [self._para(text, slot["role"], bullet) for text in texts])

    def _rect(self, slide: Any, name: str, slot: dict, cell: dict | None = None) -> None:
        box = self.d.resolve_box(slot, cell)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box.x, box.y, box.w, box.h)
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(slot["fill"])
        shape.line.fill.background()
        shape.shadow.inherit = False

    def _text(self, slide: Any, name: str, slot: dict, paras: list[_Para], cell: dict | None = None, *,
              measure: bool = True) -> None:
        box = self.d.resolve_box(slot, cell)
        shape = slide.shapes.add_textbox(box.x, box.y, box.w, box.h)
        shape.name = name
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = _ANCHOR[slot.get("anchor", "top")]
        indent = self.d.tokens["bullet"]["indent_emu"]
        for index, para in enumerate(paras):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            role = self.d.role(para.runs[0].role)
            paragraph.alignment = _ALIGN[slot.get("align", "start")]
            paragraph.line_spacing = role["line_spacing"]
            paragraph.space_after = Pt(role["space_after_pt"])
            if para.bullet is not None:
                self._bullet(paragraph, para.bullet, indent)
            for run_spec in para.runs:
                self._run(paragraph.add_run(), run_spec)
        if measure:
            needed = self._needed_height(paras, box.w, indent)
            if needed > box.h:
                self.overflows.append(Overflow(self.number, name, "height", int(needed), box.h))

    def _bullet(self, paragraph: Any, char: str, indent: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        p_pr.set("marL", str(indent))
        p_pr.set("indent", str(-indent))
        color = OxmlElement("a:buClr")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", self.d.color(self.d.tokens["bullet"]["color"]))
        color.append(srgb)
        font = OxmlElement("a:buFont")
        font.set("typeface", self.d.tokens["font"]["latin"])
        bullet = OxmlElement("a:buChar")
        bullet.set("char", char)
        for element in (color, font, bullet):
            p_pr.append(element)

    def _run(self, run: Any, spec: _Run) -> None:
        role = self.d.role(spec.role)
        run.text = spec.text
        font = run.font
        font.size = Pt(spec.size_pt or role["size_pt"])
        font.bold = role["bold"]
        font.italic = role["italic"]
        font.color.rgb = self._rgb(role["color"])
        font.name = self.d.tokens["font"]["latin"]
        r_pr = run._r.get_or_add_rPr()
        complex_script = OxmlElement("a:cs")
        complex_script.set("typeface", self.d.tokens["font"]["complex_script"])
        r_pr.find(qn("a:latin")).addnext(complex_script)

    def _needed_height(self, paras: list[_Para], width: int, indent: int) -> float:
        total = 0.0
        for index, para in enumerate(paras):
            role = self.d.role(para.runs[0].role)
            size = para.runs[0].size_pt or role["size_pt"]
            available = width * self.m.width_safety - (indent if para.bullet else 0)
            text = "".join(run.text for run in para.runs)
            lines = self.m.wrap(text, available, size, role["bold"], role["italic"])
            total += len(lines) * self.m.line_height(size, role["line_spacing"])
            if index < len(paras) - 1:
                total += role["space_after_pt"] * EMU_PER_PT
        return total

    # ---- charts -------------------------------------------------------------------------------------------------

    def _chart(self, slide: Any, box: Box, spec: ChartSpec, table: Table) -> None:
        cells = {(cell.row, cell.col): cell.text.strip() for cell in table.cells}
        rows = sorted({cell.row for cell in table.cells} - set(table.header_rows))

        def cell_text(row: int, col: int) -> str:
            text = cells.get((row, col), "")
            if not text:
                raise DeckRenderError(f"CHART_CELL_MISSING:{table.table_id}:r{row}c{col}")
            return text

        parsed = []
        data = []
        for n, col in enumerate(spec.y_columns, start=1):
            header = " ".join(cells[(r, col)] for r in table.header_rows if cells.get((r, col)))
            values = []
            for row in rows:
                text = cell_text(row, col)
                match = _NUMBER.fullmatch(text)
                if match is None:
                    raise DeckRenderError(f"CHART_CELL_NOT_NUMERIC:{table.table_id}:r{row}c{col}:{text!r}")
                parsed.append(match)
                value = Decimal(match["int"].replace(",", "") + "." + (match["frac"] or "0"))
                values.append(float(-value if match["sign"] else value))
            data.append((header or self.d.label("series_fallback", n=n), values))

        decimals = max(len(m["frac"] or "") for m in parsed)
        number_format = ("#,##0" if any("," in m["int"] for m in parsed) else "0")
        number_format += ("." + "0" * decimals if decimals else "") + ('"%"' if any(m["pct"] for m in parsed) else "")

        chart_data = CategoryChartData(number_format=number_format)
        chart_data.categories = [cell_text(row, spec.x_column) for row in rows]
        for name, values in data:
            chart_data.add_series(name, values)
        frame = slide.shapes.add_chart(_CHART_TYPES[spec.chart_type], box.x, box.y, box.w, box.h, chart_data)
        frame.name = "chart"
        self._style_chart(frame.chart, spec.chart_type, number_format)

    def _style_chart(self, chart: Any, chart_type: ChartType, number_format: str) -> None:
        style = self.d.tokens["chart"]
        text_color = self._rgb(style["text_color"])
        chart.has_title = False
        chart.font.name = self.d.tokens["font"]["latin"]
        chart.font.size = Pt(style["axis_size_pt"])
        chart.font.color.rgb = text_color
        chart.has_legend = len(chart.plots[0].series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(style["legend_size_pt"])

        plot = chart.plots[0]
        plot.vary_by_categories = False
        if hasattr(plot, "gap_width"):
            plot.gap_width = style["gap_width"]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = number_format
        labels.number_format_is_linked = False
        labels.font.size = Pt(style["label_size_pt"])
        labels.font.color.rgb = text_color
        labels.position = XL_LABEL_POSITION.ABOVE if chart_type is ChartType.LINE else XL_LABEL_POSITION.OUTSIDE_END

        for series, token in zip(plot.series, style["series"]):
            color = self._rgb(token)
            if chart_type is ChartType.LINE:
                series.smooth = False
                series.format.line.color.rgb = color
                series.format.line.width = Pt(style["line_width_pt"])
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = color
            else:
                series.invert_if_negative = False
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
                series.format.line.fill.background()

        value_axis = chart.value_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = self._rgb(style["gridline_color"])
        value_axis.major_gridlines.format.line.width = Pt(style["gridline_width_pt"])
        value_axis.format.line.fill.background()
        value_axis.major_tick_mark = XL_TICK_MARK.NONE
        value_axis.tick_labels.number_format = number_format
        value_axis.tick_labels.number_format_is_linked = False
        category_axis = chart.category_axis
        category_axis.format.line.color.rgb = self._rgb(style["gridline_color"])
        category_axis.major_tick_mark = XL_TICK_MARK.NONE
