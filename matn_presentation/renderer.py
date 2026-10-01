from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)
BG = RGBColor(247, 248, 246)
INK = RGBColor(25, 39, 48)
MUTED = RGBColor(84, 101, 108)
ACCENT = RGBColor(21, 111, 107)
ACCENT_LIGHT = RGBColor(218, 235, 231)
GOLD = RGBColor(211, 166, 77)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(204, 215, 214)
FONT = "Aptos"


class PresentationRenderError(RuntimeError):
    pass


def _set_fill(shape: Any, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _set_line(shape: Any, color: RGBColor, width: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def _textbox(slide: Any, x: float, y: float, w: float, h: float, *, name: str) -> Any:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    return shape


def _style_run(run: Any, size: float, color: RGBColor = INK, *, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_text(
    slide: Any,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    name: str,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    shape = _textbox(slide, x, y, w, h, name=name)
    shape.text_frame.vertical_anchor = valign
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size, color, bold=bold)
    return shape


def _set_native_bullet(paragraph: Any) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"):
        for node in list(p_pr.findall(tag)):
            p_pr.remove(node)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "•")
    p_pr.append(bullet)


def _add_bullet_list(
    slide: Any,
    bullets: Iterable[dict[str, Any]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 21,
    spacing: float = 10,
    name: str = "matn.body.bullets",
) -> Any:
    shape = _textbox(slide, x, y, w, h, name=name)
    frame = shape.text_frame
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.clear()
        paragraph.level = 0
        paragraph.margin_left = Pt(19)
        paragraph.indent = Pt(-10)
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.08
        _set_native_bullet(paragraph)
        run = paragraph.add_run()
        run.text = bullet["text"]
        _style_run(run, size)
    return shape


def _add_header(slide: Any, slide_plan: dict[str, Any]) -> None:
    _add_text(slide, slide_plan["title"], 0.72, 0.38, 11.65, 0.58, size=31, bold=True, name="matn.title")
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.08), Inches(0.72), Inches(0.055))
    rule.name = "matn.decorative.accent-rule"
    _set_fill(rule, ACCENT)
    rule.line.fill.background()


def _citation_text(plan: dict[str, Any], slide_plan: dict[str, Any]) -> str:
    pages = sorted({page for citation in slide_plan["citations"] for page in citation["pdf_pages"]})
    ids: list[str] = []
    for citation in slide_plan["citations"]:
        for claim_id in citation["claim_ids"]:
            if claim_id not in ids:
                ids.append(claim_id)
    if not ids:
        page_text = ", ".join(map(str, plan["source_scope"]["pdf_pages"]))
        return f"Source: {plan['source_filename']} · physical PDF pages {page_text}"
    page_text = ", ".join(str(page) for page in pages)
    return f"Source: {plan['source_filename']} · PDF p. {page_text} · Claims: {', '.join(ids)}"


def _add_footer(slide: Any, plan: dict[str, Any], slide_plan: dict[str, Any]) -> None:
    _add_text(slide, _citation_text(plan, slide_plan), 0.72, 6.86, 11.65, 0.28, size=9.5, color=MUTED, name="matn.footer.citation", valign=MSO_ANCHOR.MIDDLE)
    _add_text(slide, f"{slide_plan['slide_number']:02d}", 12.2, 6.86, 0.4, 0.28, size=10, color=ACCENT, bold=True, name="matn.footer.number", align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)


def _render_title(slide: Any, plan: dict[str, Any], slide_plan: dict[str, Any]) -> None:
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(20, 49, 55)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.84), Inches(0.8), Inches(0.13), Inches(1.28))
    accent.name = "matn.decorative.title-accent"
    _set_fill(accent, GOLD)
    accent.line.fill.background()
    _add_text(slide, "MATN", 1.22, 0.82, 2.4, 0.34, size=14, color=GOLD, bold=True, name="matn.title.brand")
    _add_text(slide, plan["deck_title"], 1.2, 1.45, 10.8, 1.15, size=44, color=WHITE, bold=True, name="matn.title.main", valign=MSO_ANCHOR.MIDDLE)
    scope = plan["source_scope"]
    page_text = "–".join(map(str, scope["pdf_pages"])) if len(scope["pdf_pages"]) > 1 else str(scope["pdf_pages"][0])
    section_text = f"Section {scope['section']} · " if scope.get("section") else ""
    _add_text(slide, f"{section_text}English presentation MVP", 1.22, 2.82, 8.4, 0.42, size=23, color=ACCENT_LIGHT, name="matn.title.subtitle")
    _add_text(slide, "Grounded PDF summaries and presentations", 1.22, 4.6, 7.8, 0.38, size=20, color=WHITE, name="matn.title.product")
    mode_text = "Validated cached demo" if plan["generation_mode"] == "validated_demo" else "Live local generation"
    _add_text(slide, f"{mode_text} · Local only · Physical PDF pages {page_text}", 1.22, 5.2, 9.8, 0.36, size=15, color=ACCENT_LIGHT, name="matn.title.scope")
    _add_text(slide, plan["source_filename"], 1.22, 6.38, 5.5, 0.3, size=11, color=RGBColor(188, 211, 209), name="matn.title.source")
    _add_text(slide, "01", 12.15, 6.38, 0.45, 0.3, size=11, color=GOLD, bold=True, name="matn.footer.number", align=PP_ALIGN.RIGHT)


def _render_standard(slide: Any, slide_plan: dict[str, Any]) -> None:
    _add_header(slide, slide_plan)
    _add_bullet_list(slide, slide_plan["bullets"], 1.04, 1.55, 11.2, 4.92, size=21.5, spacing=13)


def _render_three_columns(slide: Any, slide_plan: dict[str, Any]) -> None:
    _add_header(slide, slide_plan)
    overview, *components = slide_plan["bullets"]
    _add_text(slide, overview["text"], 0.82, 1.42, 11.7, 0.5, size=20, color=MUTED, name="matn.architecture.overview")
    headings = ["Implementer", "Scheduler", "Worker"]
    for index, (heading, bullet) in enumerate(zip(headings, components)):
        x = 0.82 + index * 4.12
        if index:
            separator = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - 0.24), Inches(2.08), Inches(0.015), Inches(3.75))
            separator.name = f"matn.decorative.column-separator-{index}"
            _set_fill(separator, LINE)
            separator.line.fill.background()
        _add_text(slide, f"0{index + 1}", x, 2.08, 0.54, 0.42, size=14, color=ACCENT, bold=True, name=f"matn.architecture.index-{index}")
        _add_text(slide, heading, x, 2.62, 3.55, 0.5, size=24, bold=True, name=f"matn.architecture.heading-{index}")
        _add_text(slide, bullet["text"], x, 3.35, 3.55, 1.6, size=20, color=MUTED, name=f"matn.architecture.body-{index}")


def _render_process(slide: Any, slide_plan: dict[str, Any]) -> None:
    _add_header(slide, slide_plan)
    for index, bullet in enumerate(slide_plan["bullets"]):
        y = 1.43 + index * 0.98
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.92), Inches(y), Inches(0.48), Inches(0.48))
        circle.name = f"matn.process.index-{index}"
        _set_fill(circle, ACCENT if index < 4 else GOLD)
        circle.line.fill.background()
        text_frame = circle.text_frame
        text_frame.clear()
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(index + 1)
        _style_run(run, 14, WHITE, bold=True)
        _add_text(slide, bullet["text"], 1.72, y - 0.02, 10.45, 0.62, size=20.5, name=f"matn.process.step-{index}", valign=MSO_ANCHOR.MIDDLE)


def _render_two_columns(slide: Any, slide_plan: dict[str, Any]) -> None:
    _add_header(slide, slide_plan)
    left = slide_plan["bullets"][:2]
    right = slide_plan["bullets"][2:]
    _add_text(slide, "Scheduler", 0.88, 1.52, 5.3, 0.5, size=24, bold=True, color=ACCENT, name="matn.columns.left-title")
    _add_bullet_list(slide, left, 0.92, 2.18, 5.25, 3.8, size=20.5, spacing=16, name="matn.columns.left-body")
    separator = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.56), Inches(1.55), Inches(0.018), Inches(4.75))
    separator.name = "matn.decorative.two-column-separator"
    _set_fill(separator, LINE)
    separator.line.fill.background()
    _add_text(slide, "Worker", 6.93, 1.52, 5.3, 0.5, size=24, bold=True, color=ACCENT, name="matn.columns.right-title")
    _add_bullet_list(slide, right, 6.97, 2.18, 5.2, 3.8, size=20.5, spacing=13, name="matn.columns.right-body")


def _add_picture_contain(slide: Any, image_path: Path, x: float, y: float, w: float, h: float) -> Any:
    with Image.open(image_path) as image:
        aspect = image.width / image.height
    box_aspect = w / h
    if aspect >= box_aspect:
        draw_w = w
        draw_h = w / aspect
        draw_x = x
        draw_y = y + (h - draw_h) / 2
    else:
        draw_h = h
        draw_w = h * aspect
        draw_x = x + (w - draw_w) / 2
        draw_y = y
    shape = slide.shapes.add_picture(str(image_path), Inches(draw_x), Inches(draw_y), width=Inches(draw_w), height=Inches(draw_h))
    shape.name = "matn.figure2.source-crop"
    return shape


def _render_image(slide: Any, slide_plan: dict[str, Any], project_root: Path) -> None:
    _add_header(slide, slide_plan)
    asset_path = project_root / slide_plan["visual_asset"]["path"]
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.38), Inches(7.05), Inches(5.05))
    frame.name = "matn.decorative.image-frame"
    _set_fill(frame, WHITE)
    _set_line(frame, LINE, 0.8)
    _add_picture_contain(slide, asset_path, 0.9, 1.54, 6.75, 4.72)
    _add_bullet_list(slide, slide_plan["bullets"], 8.18, 1.58, 4.25, 4.75, size=19.5, spacing=12, name="matn.figure2.findings")


def _render_sources(slide: Any, plan: dict[str, Any], slide_plan: dict[str, Any]) -> None:
    _add_header(slide, slide_plan)
    _add_bullet_list(slide, slide_plan["bullets"], 0.98, 1.55, 7.55, 3.3, size=21, spacing=15, name="matn.takeaways")
    _add_text(slide, "SOURCE DISCLOSURE", 9.08, 1.58, 3.25, 0.34, size=12, bold=True, color=ACCENT, name="matn.disclosure.label")
    scope = plan["source_scope"]
    pages = "–".join(map(str, scope["pdf_pages"])) if len(scope["pdf_pages"]) > 1 else str(scope["pdf_pages"][0])
    section = f"Section {scope['section']}\n" if scope.get("section") else ""
    _add_text(slide, f"{plan['source_filename']}\n{section}Physical PDF pages {pages}", 9.08, 2.18, 3.25, 1.65, size=20, bold=True, name="matn.disclosure.source")
    note = "Bullets cite accepted claim IDs. No OCR or cloud inference is used."
    if any(item.get("visual_asset") for item in plan["slides"]):
        note = "Bullets cite accepted claim IDs. Figure 2 uses a direct PDF crop and verified caption facts. No OCR or raw picture transcription is used."
    _add_text(slide, note, 9.08, 4.18, 3.25, 1.38, size=14.5, color=MUTED, name="matn.disclosure.note")


def render_presentation(plan: dict[str, Any], output_path: Path, *, project_root: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    presentation.core_properties.title = plan["deck_title"]
    presentation.core_properties.subject = "Matn Phase 04A grounded English presentation"
    presentation.core_properties.author = "Matn"
    presentation.core_properties.keywords = "grounded citations, local AI, automated research"
    presentation.core_properties.comments = "Generated deterministically from a validated Matn deck plan."
    blank = presentation.slide_layouts[6]
    for slide_plan in plan["slides"]:
        slide = presentation.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG
        layout = slide_plan["layout"]
        if layout == "title":
            _render_title(slide, plan, slide_plan)
        else:
            if layout == "three_columns":
                _render_three_columns(slide, slide_plan)
            elif layout == "process":
                _render_process(slide, slide_plan)
            elif layout == "two_columns":
                _render_two_columns(slide, slide_plan)
            elif layout == "image_and_caption":
                _render_image(slide, slide_plan, project_root)
            elif layout == "sources":
                _render_sources(slide, plan, slide_plan)
            else:
                _render_standard(slide, slide_plan)
            _add_footer(slide, plan, slide_plan)
        try:
            notes = slide.notes_slide.notes_text_frame
            notes.text = _citation_text(plan, slide_plan)
        except (AttributeError, NotImplementedError):
            pass
    presentation.save(output_path)


@dataclass(frozen=True)
class Bounds:
    left: int
    top: int
    right: int
    bottom: int


def _bounds(shape: Any) -> Bounds:
    return Bounds(shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)


def _overlap_area(a: Bounds, b: Bounds) -> int:
    width = max(0, min(a.right, b.right) - max(a.left, b.left))
    height = max(0, min(a.bottom, b.bottom) - max(a.top, b.top))
    return width * height


def validate_rendered_pptx(path: Path, *, expected_slides: int) -> dict[str, Any]:
    presentation = Presentation(path)
    if len(presentation.slides) != expected_slides:
        raise PresentationRenderError("RENDERED_SLIDE_COUNT_MISMATCH")
    if presentation.slide_width != SLIDE_W or presentation.slide_height != SLIDE_H:
        raise PresentationRenderError("RENDERED_SLIDE_SIZE_MISMATCH")
    overlap_warnings: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        visible = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > SLIDE_W or shape.top + shape.height > SLIDE_H:
                raise PresentationRenderError(f"SHAPE_OUT_OF_BOUNDS:slide={slide_index}:shape={shape.name}")
            if shape.width <= 0 or shape.height <= 0:
                raise PresentationRenderError(f"INVALID_SHAPE_SIZE:slide={slide_index}:shape={shape.name}")
            if "decorative" not in shape.name and not shape.name.startswith("matn.footer"):
                visible.append(shape)
        for left_index, left in enumerate(visible):
            for right in visible[left_index + 1:]:
                area = _overlap_area(_bounds(left), _bounds(right))
                if area <= 0:
                    continue
                smaller = min(left.width * left.height, right.width * right.height)
                if smaller and area / smaller > 0.03:
                    overlap_warnings.append(f"slide={slide_index}:{left.name}|{right.name}")
    if overlap_warnings:
        raise PresentationRenderError("VISIBLE_OBJECT_OVERLAP:" + ";".join(overlap_warnings))
    return {"status": "pass", "slide_count": len(presentation.slides), "overlaps": 0, "out_of_bounds": 0}
