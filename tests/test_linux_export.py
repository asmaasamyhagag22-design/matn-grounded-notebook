"""Exercise the actual Linux export, including multi-slide page count and PNG size."""
import os
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pypdfium2 as pdfium

from scripts.export_deck import export_deck


@pytest.mark.skipif(os.name == 'nt', reason='Linux LibreOffice backend; Windows uses PowerPoint COM')
def test_libreoffice_exports_every_slide_to_pdf_and_png(tmp_path):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for title in ['Migration smoke 1', 'Migration smoke 2']:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = 'Liberation Sans'
    pptx = tmp_path / 'deck.pptx'
    presentation.save(pptx)
    export_deck(pptx, tmp_path / 'png', tmp_path / 'deck.pdf')
    document = pdfium.PdfDocument(str(tmp_path / 'deck.pdf'))
    try:
        assert len(document) == 2
        text_page = document[0].get_textpage()
        try:
            assert 'Migration smoke 1' in text_page.get_text_range()
        finally:
            text_page.close()
    finally:
        document.close()
    assert len(list((tmp_path / 'png').glob('slide_*.png'))) == 2
    for path in (tmp_path / 'png').glob('slide_*.png'):
        with Image.open(path) as image:
            assert image.size == (1280, 720)
