from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE

from matn_contracts.domain import DeckPlan, Slide
from matn_contracts.enums import FactStatus, SlideType
from matn_deck import DeckOverflowError, DeckRenderError, build_presentation, contrast_ratio, load_design, render
from matn_deck.specimen import item, sample_fact, sample_table, specimen_inputs, SPECIMEN_DOC_ID


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def specimen():
    plan, facts, tables = specimen_inputs()
    prs, report = build_presentation(plan, facts, tables)
    return plan, facts, tables, prs, report


def shape_names(slide) -> set[str]:
    return {shape.name for shape in slide.shapes}


def shape_text(slide, name: str) -> str:
    return next(shape for shape in slide.shapes if shape.name == name).text_frame.text


def signature(prs) -> list:
    """Everything a reader sees on a slide, in order: shape name, kind, geometry, text, chart data, notes."""
    result = []
    for slide in prs.slides:
        for shape in slide.shapes:
            entry = [shape.name, int(shape.shape_type), shape.left, shape.top, shape.width, shape.height]
            if shape.has_text_frame:
                entry.append(shape.text_frame.text)
            if shape.has_chart:
                plot = shape.chart.plots[0]
                entry.append([list(plot.categories), [(s.name, list(s.values)) for s in plot.series]])
            result.append(entry)
        result.append(slide.notes_slide.notes_text_frame.text)
    return result


def test_every_archetype_renders_without_overflow(specimen, tmp_path):
    plan, facts, tables, prs, report = specimen
    assert {slide.slide_type for slide in plan.slides} == set(SlideType)
    assert len(prs.slides) == len(plan.slides)
    assert report.overflows == ()
    expected = {
        SlideType.TITLE: {"title", "subtitle"},
        SlideType.EXECUTIVE_SUMMARY: {"title", "item_1_number", "item_4_text", "footer_source", "footer_page"},
        SlideType.KPI: {"kpi_1_value", "kpi_4_label", "kpi_4_source", "items", "footer_source"},
        SlideType.FINDING: {"items", "panel_heading", "quote", "footer_source"},
        SlideType.CHART: {"chart", "items", "footer_source"},
        SlideType.IMPLICATIONS: {"item_1_text", "item_4_number", "footer_source"},
        SlideType.LIMITATIONS: {"items", "footer_source"},
        SlideType.SOURCES: {"list", "panel_body", "footer_page"},
    }
    for slide_plan, slide in zip(plan.slides, prs.slides):
        assert expected[slide_plan.slide_type] <= shape_names(slide), slide_plan.slide_type
        assert shape_text(slide, "title") == slide_plan.title
    sources = prs.slides[-1]
    assert "footer_source" not in shape_names(sources)

    out = tmp_path / "specimen.pptx"
    assert render(plan, facts, tables, out, strict=True) == report
    assert len(Presentation(out).slides) == len(plan.slides)


def test_all_text_is_editable_and_there_are_no_pictures(specimen):
    plan, _, _, prs, _ = specimen
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.shape_type in {MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.CHART}
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                assert shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.RECTANGLE
                assert shape.text_frame.text == ""
    texts = {slide_plan.title: slide for slide_plan, slide in zip(plan.slides, prs.slides)}
    for slide_plan in plan.slides:
        all_text = "\n".join(s.text_frame.text for s in texts[slide_plan.title].shapes if s.has_text_frame)
        for slide_item in slide_plan.items:
            assert slide_item.text in all_text


def test_contract_limits_are_respected_by_layouts_and_renderer(specimen):
    layouts = load_design().layouts["archetypes"]
    max_items = next(m.max_length for m in Slide.model_fields["items"].metadata if hasattr(m, "max_length"))
    assert max_items == 4
    assert len(layouts["executive_summary"]["repeat"]["cells"]) == max_items
    assert len(layouts["implications"]["repeat"]["cells"]) == max_items
    assert set(layouts["kpi"]["repeat"]["cells_by_count"]) == {"2", "3", "4"}
    assert set(layouts) == {t.value for t in SlideType}

    plan, facts, tables, _, _ = specimen
    rejected = sample_fact("Sample: a rejected claim.", "Sample text: a rejected claim.", 1,
                           status=FactStatus.REJECTED)
    bad = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
        Slide(slide_type=SlideType.LIMITATIONS, title="Sample assertion", items=(item("Sample.", rejected),),
              notes=None),))
    with pytest.raises(DeckRenderError, match="REJECTED_FACT_CITED"):
        build_presentation(bad, facts + (rejected,), tables)
    with pytest.raises(ValueError, match="UNKNOWN_FACT_ID"):
        build_presentation(bad, facts, tables)


def test_overflow_is_detected_on_an_overlong_sample():
    facts = (sample_fact("Sample: wide.", "Sample text: wide.", 0),)
    wide = "W" * 120
    plan = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
        Slide(slide_type=SlideType.FINDING, title="W" * 90, items=tuple(item(wide, facts[0]) for _ in range(4)),
              notes=None),))
    prs, report = build_presentation(plan, facts, ())
    slots = {(o.slide_number, o.slot) for o in report.overflows}
    assert {(1, "title"), (1, "items")} <= slots
    assert all(o.needed_emu > o.available_emu for o in report.overflows)
    assert shape_text(prs.slides[0], "title") == "W" * 90  # reported, never cut or auto-shrunk
    with pytest.raises(DeckOverflowError):
        build_presentation(plan, facts, (), strict=True)


def test_colors_fonts_and_sizes_come_from_tokens(specimen):
    tokens = load_design().tokens
    colors = set(tokens["color"].values())
    fonts = {tokens["font"]["latin"], tokens["font"]["complex_script"]}
    sizes = {role["size_pt"] for role in tokens["type"].values()} | set(tokens["kpi_value_steps_pt"])
    sizes |= {tokens["chart"][k] for k in ("axis_size_pt", "label_size_pt", "legend_size_pt")}
    _, _, _, prs, _ = specimen
    for slide in prs.slides:
        parts = [slide.part] + [rel.target_part for rel in slide.part.rels.values() if "chart" in rel.reltype]
        for part in parts:
            xml = part.blob.decode("utf-8") if hasattr(part, "blob") else ""
            assert set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)) <= colors
            assert set(re.findall(r'typeface="([^"]+)"', xml)) <= fonts
            assert {int(v) / 100 for v in re.findall(r'\ssz="(\d+)"', xml)} <= sizes
        assert slide.background.fill.fore_color.rgb is not None
    source = "\n".join(p.read_text(encoding="utf-8") for p in (ROOT / "matn_deck").glob("*.py"))
    assert not re.search(r"RGBColor\(0x|RGBColor\.from_string\(\"|['\"][0-9A-Fa-f]{6}['\"]", source)
    assert not re.search(r"Calibri|Segoe|Carlito|Pt\(\d|Inches\(|Emu\(\d", source)


def test_the_same_plan_gives_the_same_slide_content():
    plan, facts, tables = specimen_inputs()
    first, _ = build_presentation(plan, facts, tables)
    second, _ = build_presentation(plan, facts, tables)
    assert signature(first) == signature(second)


def test_kpi_values_and_chart_data_are_copied_from_facts_and_tables(specimen):
    plan, facts, tables, prs, _ = specimen
    by_id = {f.fact_id: f for f in facts}
    kpi_index = next(i for i, s in enumerate(plan.slides) if s.slide_type is SlideType.KPI)
    for n, fact_id in enumerate(plan.slides[kpi_index].kpi_fact_ids, start=1):
        value = by_id[fact_id].number_value
        text = shape_text(prs.slides[kpi_index], f"kpi_{n}_value")
        assert text.startswith(value.value_as_written)
        assert value.unit is None or text.endswith(value.unit)
        assert shape_text(prs.slides[kpi_index], f"kpi_{n}_label") == value.label

    chart_index = next(i for i, s in enumerate(plan.slides) if s.slide_type is SlideType.CHART)
    chart = next(s for s in prs.slides[chart_index].shapes if s.has_chart).chart
    plot = chart.plots[0]
    assert list(plot.categories) == ["Region A", "Region B", "Region C", "Region D"]
    assert [s.name for s in plot.series] == ["2024", "2025"]
    assert list(plot.series[0].values) == [12.5, 9.0, 6.5, 4.0]
    assert "p. 10 (table)" in shape_text(prs.slides[chart_index], "footer_source")

    broken = sample_table([["Region", "Value"], ["Region A", "about 12"]], 9)
    spec = plan.slides[chart_index].chart.model_copy(update={"table_id": broken.table_id, "y_columns": (1,)})
    bad = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(plan.slides[chart_index].model_copy(update={"chart": spec}),))
    with pytest.raises(DeckRenderError, match="CHART_CELL_NOT_NUMERIC"):
        build_presentation(bad, facts, (broken,))


def test_every_content_slide_names_its_source_pages(specimen):
    plan, _, _, prs, _ = specimen
    for slide_plan, slide in zip(plan.slides, prs.slides):
        if slide_plan.slide_type in {SlideType.TITLE, SlideType.SOURCES}:
            continue
        assert re.fullmatch(r"Source: p\. \d+.*", shape_text(slide, "footer_source"))
        assert "fact:" in slide.notes_slide.notes_text_frame.text
    assert "Verified facts cited:" in shape_text(prs.slides[-1], "panel_body")


def test_every_text_role_meets_the_contrast_minimum():
    tokens = load_design().tokens
    for name, role in tokens["type"].items():
        ratio = contrast_ratio(tokens["color"][role["color"]], tokens["color"][role["on"]])
        assert ratio >= tokens["contrast_min"], (name, round(ratio, 2))
