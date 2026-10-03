"""Template contract, validator and gallery (docs/design/template-contract.md)."""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from matn_deck.colors import contrast, hex_rgb
from matn_deck.templates import (
    GALLERY_DIR,
    LAYOUTS,
    default_template_id,
    gallery,
    is_standard_font,
    validate_template,
)

ROOT = Path(__file__).resolve().parents[1]
GALLERY_SPEC = json.loads((GALLERY_DIR / "src" / "gallery.json").read_text(encoding="utf-8"))
CHOSEN = [entry["id"] for entry in GALLERY_SPEC["templates"] if not entry.get("excluded")]


def _expand(cell: str) -> tuple[str, ...]:
    """"item_1..item_4, source" or "kpi_value_1..2" -> placeholder names."""
    names: list[str] = []
    for part in filter(None, (p.strip() for p in cell.split(","))):
        match = re.fullmatch(r"([a-z_]+)_(\d)(_title)?\.\.(?:[a-z_]+_)?(\d)(?:_title)?", part)
        if match is None:
            names.append(part)
            continue
        stem, first, suffix, last = match.group(1), int(match.group(2)), match.group(3) or "", int(match.group(4))
        names += [f"{stem}_{i}{suffix}" for i in range(first, last + 1)]
    return tuple(names)


def test_contract_document_and_code_name_the_same_layouts_and_placeholders():
    text = (ROOT / "docs" / "design" / "template-contract.md").read_text(encoding="utf-8")
    rows = re.findall(r"^\| `(matn:[^`]+)` \|([^|]*)\|([^|]*)\|\s*$", text, re.M)
    documented = {name: (set(_expand(req)), set(_expand(opt))) for name, req, opt in rows}
    assert documented == {name: (set(req), set(opt)) for name, (req, opt) in LAYOUTS.items()}


def test_every_gallery_template_passes_and_carries_its_metadata():
    found = {template.id: template for template in gallery()}
    assert sorted(found) == sorted(CHOSEN)
    assert default_template_id() in found
    for template_id, template in found.items():
        report = validate_template(template.path)
        assert report.ok and report.warnings == (), (template_id, report)
        assert template.path.stat().st_size < 5 * 1024 * 1024
        assert template.thumbnail.is_file()
        meta = template.meta
        required = {"id", "name_ar", "name_en", "fonts", "mode", "source", "content_area_pt", "divider",
                    "pattern", "glow", "family", "mood", "audience"}
        assert required <= set(meta)
        assert meta["id"] == template_id and meta["mode"] in {"dark", "light"}
        assert meta["family"] in {"classic", "creative"} and meta["mood"] and meta["audience"]
        if meta["family"] == "classic":
            assert meta["source"].startswith("Office theme ")
        else:
            assert meta["source"] == f"design/creative/{template_id}/spec.json"
            assert {"roles", "card_style", "divider_style", "background"} <= set(meta)
        assert all(is_standard_font(face) for face in meta["fonts"].values())


def test_variant_geometry_covers_exactly_the_contract_layouts():
    from matn_deck.variants import load_variants
    assert set(load_variants()["variants"]) == set(LAYOUTS)


def _layout(template_path: Path, name: str):
    return next(layout for layout in Presentation(str(template_path)).slide_layouts if layout.name == name)


def test_dividers_are_full_bleed_theme_color_with_checked_contrast():
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for template in gallery():
        layout = _layout(template.path, "matn:section_divider:default")
        bg = layout._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}bgPr")
        fills = [hex_rgb(c.get("val")) for c in bg.iter(f"{a}srgbClr")]
        assert fills, template.id
        texts = {}
        for ph in layout.placeholders:
            if ph.name not in ("title", "section_number", "subtitle"):
                continue
            color = ph._element.find(f".//{a}defRPr/{a}solidFill/{a}srgbClr")
            if color is None:
                color = ph._element.find(f".//{a}defRPr/{a}ln/{a}solidFill/{a}srgbClr")
            assert color is not None, (template.id, ph.name)
            texts[ph.name] = hex_rgb(color.get("val"))
        assert set(texts) == {"title", "section_number", "subtitle"}, template.id
        worst = min(contrast(text, fill) for text in texts.values() for fill in fills)
        assert worst >= 4.5, (template.id, worst)
        assert template.meta["divider"]["contrast"] >= 4.5
        number = next(ph for ph in layout.placeholders if ph.name == "section_number")
        size = int(number._element.find(f".//{a}defRPr").get("sz")) / 100
        assert 160 <= size <= 200, (template.id, size)


def test_pattern_and_glow_follow_the_theme_flags():
    for template in gallery():
        divider = _layout(template.path, "matn:section_divider:default")
        content = _layout(template.path, "matn:finding:cards-3")
        patterned = [s for s in divider.shapes if s.name == "pattern"]
        assert bool(patterned) == bool(template.meta["pattern"]), template.id
        assert not [s for s in content.shapes if s.name == "pattern"]
        glow = [s for s in content.shapes if s.name == "glow"]
        assert bool(glow) == template.meta["glow"], template.id
        # no layout carries card shapes any more: the renderer draws cards for filled slots only
        assert not [s for layout in Presentation(str(template.path)).slide_layouts for s in layout.shapes
                    if s.name == "card"]


def _broken_copy(tmp_path: Path, mutate) -> Path:
    folder = tmp_path / "broken"
    shutil.copytree(GALLERY_DIR / default_template_id(), folder)
    prs = Presentation(str(folder / "template.pptx"))
    mutate(prs)
    prs.save(str(folder / "template.pptx"))
    return folder


def test_validator_fails_with_a_clear_reason_when_a_required_placeholder_is_missing(tmp_path):
    def drop_quote(prs):
        layout = next(layout for layout in prs.slide_layouts if layout.name == "matn:finding:split-60-40")
        quote = next(ph for ph in layout.placeholders if ph.name == "quote")
        quote._element.getparent().remove(quote._element)

    folder = _broken_copy(tmp_path, drop_quote)
    report = validate_template(folder / "template.pptx")
    assert not report.ok
    assert report.errors == ("PLACEHOLDER_MISSING:matn:finding:split-60-40:quote",)
    assert gallery(tmp_path) == ()  # a failing template never reaches the gallery


def test_validator_rejects_sample_slides_wrong_size_and_odd_fonts(tmp_path):
    def spoil(prs):
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.slide_height = prs.slide_width * 3 // 4
        layout = prs.slide_layouts[1]
        layout.placeholders[0].text_frame.text = "x"
        layout.placeholders[0].text_frame.paragraphs[0].runs[0].font.name = "Comic Sans MS"

    report = validate_template(_broken_copy(tmp_path, spoil) / "template.pptx")
    assert {e.split(":")[0] for e in report.errors} == {"SAMPLE_SLIDES_LEFT", "SLIDE_SIZE_NOT_16_9",
                                                        "FONT_NOT_STANDARD"}
    assert validate_template(tmp_path / "missing.pptx").errors == ("TEMPLATE_UNREADABLE:PackageNotFoundError",)


def test_standard_font_names():
    for face in ("Aptos Display", "Bahnschrift SemiBold", "Segoe UI Semibold", "Calibri Light", "Sakkal Majalla"):
        assert is_standard_font(face), face
    for face in ("Tw Cen MT", "Comic Sans MS", "Segoe UI Fancy Script"):
        assert not is_standard_font(face), face
