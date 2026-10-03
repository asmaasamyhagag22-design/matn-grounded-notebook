"""Creative-family color, spec and deterministic-background behavior."""
from __future__ import annotations

import hashlib
import json
import shutil

import pytest
from pydantic import ValidationError
from pptx import Presentation

from design.creative.backgrounds import content_zone_variance, decor_zone_metrics, render_background
from design.creative.generate import specs
from design.creative.spec import CreativeSpec, load_spec
from matn_deck.color import blend, contrast_ratio, derive_roles, oklch_to_srgb, rgb_from_hex, srgb_to_oklch
from matn_deck.decor import CARD_ALPHA
from matn_deck.specimen import specimen_inputs
from matn_deck.template_renderer import build_presentation
from matn_deck.templates import GALLERY_DIR, default_template_id

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def sample_spec(**changes):
    data = {
        "id": "sample", "name_ar": "عينة", "name_en": "Sample", "mood": "Calm", "recommended_audience": "Leaders",
        "family": "creative", "mode": "dark",
        "seed_colors": {"primary": "#123B66", "accent": "#D6A84B", "surface": "#081B33"},
        "fonts": {"title": "Bahnschrift", "body": "Segoe UI"},
        "background": {"generator": "aurora", "parameters": {"grain": 0.5}, "fixed_seed": 7},
        "card_style": {"fill_alpha": 0.16, "corner_radius_pt": 14,
                       "shadow": {"enabled": True, "alpha": 0.16, "blur_pt": 12, "distance_pt": 2}},
        "divider_style": {"type": "gradient", "secondary": "#087E8B", "thin_rules": False},
        "motif": {"type": "aurora", "placement": "edges", "opacity": 0.12}, "pattern": False,
    }
    data.update(changes)
    return data


def test_oklch_round_trip_and_known_wcag_pairs():
    for value in ("000000", "FFFFFF", "0057B8", "D6A84B", "F4EBDD"):
        rgb = rgb_from_hex(value)
        assert oklch_to_srgb(srgb_to_oklch(rgb)) == pytest.approx(rgb, abs=2e-6)
    assert contrast_ratio(rgb_from_hex("000000"), rgb_from_hex("FFFFFF")) == pytest.approx(21)
    assert contrast_ratio(rgb_from_hex("777777"), rgb_from_hex("FFFFFF")) == pytest.approx(4.478, abs=0.002)


@pytest.mark.parametrize("mode", ["light", "dark"])
def test_derived_roles_meet_contrast_targets(mode):
    roles = {name: rgb_from_hex(value) for name, value in derive_roles("#075985", "#D6A84B", mode,
                                                                        card_alpha=0.18).items()}
    for role in ("on_surface", "muted", "primary"):
        assert contrast_ratio(roles[role], roles["surface"]) >= 4.5
    assert contrast_ratio(roles["accent"], roles["surface"]) >= 3
    assert contrast_ratio(roles["focus"], roles["surface"]) >= 3
    assert contrast_ratio(roles["on_divider"], roles["divider_background"]) >= 4.5
    glass = blend(roles["on_surface"], roles["surface"], 0.18)
    assert contrast_ratio(roles["on_surface"], glass) >= 4.5
    if mode == "light":
        assert roles["accent"] == roles["primary"]  # vivid seed is a fill, never saffron text on light


def test_spec_rejects_nonstandard_font_and_missing_seed_role():
    bad_font = sample_spec(fonts={"title": "Comic Sans MS", "body": "Segoe UI"})
    with pytest.raises(ValidationError, match="FONT_NOT_STANDARD"):
        CreativeSpec.model_validate(bad_font)
    missing = sample_spec()
    del missing["seed_colors"]["accent"]
    with pytest.raises(ValidationError, match="accent"):
        CreativeSpec.model_validate(missing)


def test_all_nine_specs_validate_and_gallery_metadata_matches():
    found = specs()
    assert len(found) == 9
    gallery = json.loads((GALLERY_DIR / "src" / "gallery.json").read_text(encoding="utf-8"))["templates"]
    creative = {entry["id"]: entry for entry in gallery if entry["family"] == "creative"}
    assert set(creative) == set(found)
    for template_id, (path, spec) in found.items():
        assert load_spec(path) == spec
        assert creative[template_id]["mood"] == spec.mood
        assert creative[template_id]["audience"] == spec.recommended_audience
    source = (GALLERY_DIR.parent / "creative" / "generate.py").read_text(encoding="utf-8")
    variations = set(found) - {"nile-night", "papyrus", "mashrabiya"}
    assert not any(template_id in source for template_id in variations)  # variations add data, never a code branch


def test_background_art_is_perceptible_while_content_stays_calm():
    for template_id, (_, spec) in specs().items():
        image = render_background(spec, (640, 360))
        safe_left = 0.2
        if spec.background.generator == "mashrabiya":
            safe_left = float(spec.background.parameters["band_width"]) + float(spec.background.parameters["gutter"])
        metrics = decor_zone_metrics(image, (safe_left, 0.18, 0.84 - safe_left, 0.66))
        assert metrics["calm"], (template_id, metrics)
        assert metrics["perceptible"], (template_id, metrics)


@pytest.mark.parametrize("kind", ["aurora", "paper", "mashrabiya"])
def test_backgrounds_are_deterministic_and_keep_the_content_zone_calm(kind):
    data = sample_spec(mode="light" if kind != "aurora" else "dark")
    data["background"] = {"generator": kind, "parameters": {}, "fixed_seed": 91}
    data["motif"] = {"type": "lotus" if kind == "paper" else "mashrabiya" if kind == "mashrabiya" else "aurora",
                     "placement": "corners" if kind == "paper" else "side_band" if kind == "mashrabiya" else "edges",
                     "opacity": 0.12}
    spec = CreativeSpec.model_validate(data)
    first, second = render_background(spec, (320, 180)), render_background(spec, (320, 180))
    assert hashlib.sha256(first.tobytes()).digest() == hashlib.sha256(second.tobytes()).digest()
    assert content_zone_variance(first) < 12


def _alpha(shape) -> int:
    return int(shape._element.find(f".//{_A}solidFill/{_A}schemeClr/{_A}alpha").get("val"))


def test_card_and_divider_tokens_are_honored_and_classic_defaults_stay_unchanged(tmp_path):
    source = GALLERY_DIR / default_template_id()
    classic_plan, facts, tables = specimen_inputs()
    classic, _ = build_presentation(classic_plan, facts, tables, source / "template.pptx", strict=True)
    classic_card = next(s for slide in classic.slides for s in slide.shapes if s.name == "card_1")
    assert _alpha(classic_card) == round(CARD_ALPHA * 100000)
    assert classic_card._element.find(f".//{_A}outerShdw") is None

    creative = tmp_path / "creative"
    creative.mkdir()
    shutil.copy2(source / "template.pptx", creative / "template.pptx")
    meta = json.loads((source / "theme.json").read_text(encoding="utf-8"))
    meta["card_style"] = {"fill_alpha": 0.19, "corner_radius_pt": 18,
                          "shadow": {"enabled": True, "alpha": 0.2, "blur_pt": 14, "distance_pt": 3}}
    meta["divider_style"] = {"type": "outline_number", "outline_color": "D6A84B"}
    (creative / "theme.json").write_text(json.dumps(meta), encoding="utf-8")
    rendered, _ = build_presentation(classic_plan, facts, tables, creative / "template.pptx", strict=True)
    card = next(s for slide in rendered.slides for s in slide.shapes if s.name == "card_1")
    assert _alpha(card) == 19000
    assert card._element.find(f".//{_A}outerShdw") is not None
    divider = next(slide for slide in rendered.slides if slide.slide_layout.name.startswith("matn:section_divider"))
    number = next(s for s in divider.placeholders if s.name == "section_number")
    run = number.text_frame.paragraphs[0].runs[0]
    assert run.font._rPr.find(f"{_A}noFill") is not None
    line = run.font._rPr.find(f"{_A}ln")
    assert int(line.get("w")) >= 4 * 12700
    assert line.find(f"{_A}solidFill/{_A}srgbClr").get("val") == "D6A84B"


@pytest.mark.parametrize("template_id", ["mashrabiya", "mashrabiya-cobalt", "mashrabiya-plum"])
def test_mashrabiya_layouts_keep_all_text_outside_the_side_band_and_gutter(template_id):
    folder = GALLERY_DIR / template_id
    meta = json.loads((folder / "theme.json").read_text(encoding="utf-8"))
    safe_left = meta["decor_zone"]["safe_left_pt"]
    assert safe_left > meta["decor_zone"]["band_right_pt"]
    assert meta["content_area_pt"][0] >= safe_left
    prs = Presentation(str(folder / "template.pptx"))
    for layout in prs.slide_layouts:
        if layout.name.startswith("matn:section_divider:"):
            continue
        for shape in layout.placeholders:
            assert shape.left / 12700 >= safe_left - 0.1, (template_id, layout.name, shape.name, shape.left / 12700)


@pytest.mark.parametrize("template_id", ["papyrus", "papyrus-lapis", "papyrus-olive"])
def test_papyrus_keeps_a_visible_accessible_outline_number(template_id):
    meta = json.loads((GALLERY_DIR / template_id / "theme.json").read_text(encoding="utf-8"))
    style = meta["divider_style"]
    assert style["type"] == "outline_number" and style["outline_width_pt"] >= 4
    outline = rgb_from_hex(style["outline_color"])
    assert min(contrast_ratio(outline, rgb_from_hex(fill)) for fill in meta["divider"]["fill"]) >= 3


def test_mashrabiya_uses_deep_text_accents_and_reserves_saffron_for_focus_fills():
    for template_id in ("mashrabiya", "mashrabiya-cobalt", "mashrabiya-plum"):
        roles = json.loads((GALLERY_DIR / template_id / "theme.json").read_text(encoding="utf-8"))["roles"]
        assert roles["accent"] == roles["primary"]
        assert roles["focus"] != roles["accent"]
