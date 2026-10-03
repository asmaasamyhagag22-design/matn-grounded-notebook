"""Demo dashboard (app/dashboard.py) and its data layer (matn_deck/demo_data.py)."""
from __future__ import annotations

import re
from pathlib import Path

import pytest

from matn_deck import demo_data

ROOT = Path(__file__).resolve().parents[1]
DASHBOARD = ROOT / "app" / "dashboard.py"
PAGES = ["الرئيسية", "معرض القوالب", "اعمل عرض", "نفس الحقايق، أي قالب", "الإيصالات", "الجودة", "الطريق"]


def test_gallery_lists_dark_templates_first_and_marks_modulo_default():
    entries = demo_data.gallery_entries()
    modes = [e["mode"] for e in entries]
    assert modes == sorted(modes, key=lambda m: m != "dark")
    defaults = [e["id"] for e in entries if e["default"]]
    assert defaults == ["modulo"]
    assert "slice" not in {e["id"] for e in entries}  # excluded by QA stays out
    assert {"slice"} <= {x["id"] for x in demo_data.excluded_templates()}


def test_scores_table_parses_the_first_table(tmp_path, monkeypatch):
    md = tmp_path / "scores.md"
    md.write_text("# t\n\n| Template | Beauty | Gate |\n|---|---|---|\n| a | 0.91 | pass |\n| b | 0.92 | pass |\n\n"
                  "| Template | Other |\n|---|---|\n| z | 1 |\n", encoding="utf-8")
    monkeypatch.setattr(demo_data, "SCORES_MD", md)
    assert demo_data.scores_table() == [{"Template": "a", "Beauty": "0.91", "Gate": "pass"},
                                        {"Template": "b", "Beauty": "0.92", "Gate": "pass"}]


def test_last_test_run_falls_back_to_the_handoff(tmp_path, monkeypatch):
    handoff = tmp_path / "HANDOFF.md"
    handoff.write_text("old `410 passed in 259.39s`\nnew `428 passed, 6 skipped in 299.50s`\n", encoding="utf-8")
    monkeypatch.setattr(demo_data, "TEST_RUN", tmp_path / "missing.json")
    monkeypatch.setattr(demo_data, "HANDOFF", handoff)
    run = demo_data.last_test_run()
    assert (run["passed"], run["skipped"], run["seconds"]) == (428, 6, 299.5)


def test_dashboard_has_no_external_urls_or_pdf_library():
    source = DASHBOARD.read_text(encoding="utf-8") + (ROOT / "matn_deck" / "demo_data.py").read_text(encoding="utf-8")
    urls = re.findall(r"https?://[^\s\"')]+", source)
    assert all(u.startswith(("http://127.0.0.1", "http://localhost")) for u in urls), urls
    assert not re.search(r"^\s*(import|from)\s+(fitz|pymupdf)", source, re.M)
    assert not re.search(r"fonts\.googleapis|//cdn|cdn\.\w|unpkg\.com|jsdelivr\.net", source, re.I)


@pytest.mark.skipif(not (demo_data.OUT / "manifest.json").is_file(), reason="demo cache not built (out/demo/)")
def test_every_dashboard_page_loads_from_the_cache():
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file(str(DASHBOARD), default_timeout=180)
    app.run()
    assert not app.exception
    for page in PAGES:
        app.sidebar.radio(key="page").set_value(page).run()
        assert not app.exception, (page, [e.value for e in app.exception])


def test_deck_numbers_compare_claims_not_the_layout_only_quote_box(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, text in (("item_1", "Growth reached 9.8% (1/2)"), ("quote", "“since Q4-FY2020/21 it rose 99 points”"),
                       ("source", "p. 7, 8"), ("marker_1", "1")):
        box = slide.shapes.add_textbox(Pt(10), Pt(10), Pt(200), Pt(40))
        box.name = name
        box.text_frame.text = text
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    assert demo_data.deck_numbers(path) == ["9.8"]
