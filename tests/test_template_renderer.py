"""Renderer v2: layouts chosen by generate-and-score, placeholders filled by name, visuals drawn for filled slots,
measured fallback and split (Step 2c)."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from matn_contracts.domain import ChartSpec, DeckPlan, Slide
from matn_contracts.enums import ChartType, FactKind, FactStatus, SlideType
from matn_deck.scoring import Rect
from matn_deck.specimen import SPECIMEN_DOC_ID, item, sample_fact, sample_table, specimen_inputs
from matn_deck.template_renderer import (
    CONTENT_TYPES,
    DeckRenderError,
    TemplateInvalidError,
    _Content,
    _DeckBuilder,
    build_presentation,
    load_labels,
    render,
)
from matn_deck.templates import GALLERY_DIR, LAYOUTS, default_template_id, gallery

DEFAULT = GALLERY_DIR / default_template_id() / "template.pptx"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.fixture(scope="module")
def specimen():
    plan, facts, tables = specimen_inputs()
    prs, report = build_presentation(plan, facts, tables, DEFAULT, strict=True)
    return plan, facts, tables, prs, report


def by_name(slide, name):
    return next(shape for shape in slide.shapes if shape.name == name)


def candidates(report):
    return {c["layout"] for entry in report.scores for c in entry["candidates"]}


def fill_slot(shape) -> str | None:
    """The theme slot of an autoshape's solid fill (accent1, tx1, ...)."""
    color = shape._element.spPr.find(f"{_A}solidFill/{_A}schemeClr")
    return None if color is None else color.get("val")


def test_specimen_exercises_every_variant_and_slide_type(specimen, tmp_path):
    plan, facts, tables, prs, report = specimen
    assert {slide.slide_type for slide in plan.slides} == set(SlideType)
    # The 1-3 item implications variants are split fallbacks (test_split_slides_never_show_an_empty_card).
    fallbacks = {"matn:implications:single-hero", "matn:implications:cards-3", "matn:implications:two-split"}
    assert set(LAYOUTS) - fallbacks <= candidates(report)
    chosen = set(report.layouts)
    for name in ("matn:kpi:hero-number", "matn:finding:single-hero", "matn:kpi:pictogram",
                 "matn:closing:takeaways", "matn:limitations:numbered", "matn:chart:hero"):
        assert name in chosen, name
    assert report.overflows == ()
    assert all(w.startswith("CHART_TYPE_BY_RULE:") for w in report.warnings), report.warnings
    assert len(prs.slides) == len(report.layouts) == len(report.scores)
    out = tmp_path / "deck.pptx"
    assert render(plan, facts, tables, DEFAULT, out, strict=True) == report


def test_every_slide_records_its_candidates_and_scores(specimen):
    _, _, _, _, report = specimen
    for entry in report.scores:
        assert entry["chosen"] in entry["candidates"] or entry["part"] is not None
        if entry["slide_type"] in {t.value for t in CONTENT_TYPES}:
            assert entry["chosen"]["hard"] == []
            assert set(entry["chosen"]["soft"]) >= {"fill", "balance", "vertical", "hierarchy", "density", "variety"}
            assert {c["step"] for c in entry["candidates"]} == {"L", "M", "S"}
            assert {c["anchor"] for c in entry["candidates"]} == {"top", "center"}


def test_selection_is_deterministic():
    plan, facts, tables = specimen_inputs()
    first = build_presentation(plan, facts, tables, DEFAULT)[1]
    second = build_presentation(plan, facts, tables, DEFAULT)[1]
    assert first.layouts == second.layouts
    assert [e["chosen"] for e in first.scores] == [e["chosen"] for e in second.scores]


def test_placeholders_are_filled_by_name_and_nothing_is_left_empty(specimen):
    plan, facts, _, prs, report = specimen
    by_id = {fact.fact_id: fact for fact in facts}
    kpi = next(e["slide"] - 1 for e in report.scores if e["plan_index"] == 3)
    slide = prs.slides[kpi]
    assert slide.slide_layout.name.startswith("matn:kpi:")
    assert by_name(slide, "title").text_frame.text.replace("\v", " ") == plan.slides[3].title
    for n, fact_id in enumerate(plan.slides[3].kpi_fact_ids, start=1):
        value = by_id[fact_id].number_value
        runs = by_name(slide, f"kpi_value_{n}").text_frame.paragraphs[0].runs
        assert runs[0].text == value.value_as_written
        assert "".join(r.text for r in runs[1:]).strip() == (value.unit or "")
        assert by_name(slide, f"kpi_label_{n}").text_frame.text == value.label
    assert by_name(slide, "source").text_frame.text == "p. 3, 5, 6, 7"
    for slide in prs.slides:
        for shape in slide.placeholders:
            assert shape.has_chart if shape.name == "chart" else shape.text_frame.text.strip(), shape.name
            assert not re.fullmatch(r".* \d+", shape.name), shape.name  # renamed to the contract name


def test_the_chart_follows_the_chart_rules(specimen):
    plan, _, tables, prs, report = specimen
    index = next(e["slide"] - 1 for e in report.scores if e["plan_index"] == 10)
    frame = by_name(prs.slides[index], "chart")
    assert frame.is_placeholder and frame.placeholder_format.type == PP_PLACEHOLDER.CHART
    chart = frame.chart
    plot = chart.plots[0]
    assert chart.chart_type is not None and "BAR" in str(chart.chart_type)  # 4 categories: horizontal bars, sorted
    assert list(plot.categories) == ["Region A", "Region B", "Region C", "Region D"]
    assert [(s.name, list(s.values)) for s in plot.series] == [("2024", [12.5, 9.0, 6.5, 4.0]),
                                                              ("2025", [18.0, 14.5, 8.0, 7.5])]
    assert by_name(prs.slides[index], "source").text_frame.text == "p. 8, 10 (table)"
    assert chart.value_axis.has_major_gridlines is False and plot.has_data_labels and chart.has_legend
    callout = plot.series[1].points[0].data_label.text_frame.text  # the maximum of the key series
    assert callout == load_labels()["callout_max"].format(value="18.0")

    years = next(e["slide"] - 1 for e in report.scores if e["plan_index"] == 6)
    chart = by_name(prs.slides[years], "chart").chart
    assert "COLUMN" in str(chart.chart_type) and not chart.has_legend  # 5 time points: columns, one series
    series = chart.plots[0].series[0]
    fills = [series.points[p].format.fill._xPr.find(f"{_A}solidFill")[0] for p in range(5)]
    accent = [f.get("val") == "accent1" for f in fills]
    assert accent == [False, False, False, False, True]  # the cited 18.0 (2025) is the key point
    assert series.points[4].data_label.text_frame.text.startswith("18.0")
    assert "CHART_TYPE_BY_RULE" in " ".join(report.warnings)


def test_text_stays_editable_and_visuals_are_native_shapes(specimen):
    _, _, _, prs, _ = specimen
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE, shape.name
            if shape.is_placeholder:
                continue
            assert shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE, shape.name
            assert not shape.text_frame.text, shape.name  # all text lives in named placeholders


def test_evidence_rule_turns_text_only_slides_into_visual_variants(specimen):
    _, _, _, prs, report = specimen
    limitations = next(e for e in report.scores if e["slide_type"] == "limitations")
    plain = [c for c in limitations["candidates"] if c["layout"] == "matn:limitations:default"]
    assert plain and all("NO_VISUAL" in c["hard"] for c in plain)
    assert limitations["chosen"]["layout"] == "matn:limitations:numbered"
    visual_names = re.compile(r"(card|bar_track|icon|dot|hero_card)_")
    for entry in report.scores:
        if entry["slide_type"] not in {t.value for t in CONTENT_TYPES}:
            continue
        slide = prs.slides[entry["slide"] - 1]
        visuals = [s for s in slide.shapes if visual_names.match(s.name + "_") or s.name.startswith(("card", "hero"))]
        assert visuals or any(s.name == "chart" for s in slide.shapes) or \
            slide.slide_layout.name.startswith("matn:kpi:"), entry["slide"]


def test_progress_bars_and_pictograms_appear_only_under_their_data_rules(specimen):
    plan, facts, tables, prs, report = specimen
    by_plan = {}
    for entry in report.scores:
        by_plan.setdefault(entry["plan_index"], []).append(entry)
    names = lambda i: {c["layout"] for e in by_plan[i] for c in e["candidates"]}  # noqa: E731
    assert "matn:kpi:progress" not in names(3)  # 1,240 and 4.5 days are not percentages
    assert "matn:kpi:progress" in names(5)  # 18% and 72%
    assert "matn:kpi:pictogram" in names(8) and "matn:kpi:pictogram" not in names(5)  # 72% is not a multiple of 10
    slide = prs.slides[by_plan[8][0]["slide"] - 1]
    icons = [s for s in slide.shapes if s.name.startswith("icon_")]
    on = lambda k: [fill_slot(s) == "accent1" for s in icons if s.name.startswith(f"icon_{k}_")]  # noqa: E731
    assert on(1) == [True] * 3 + [False] * 7  # "3 in 10"
    assert on(2) == [True] * 4 + [False] * 6  # 40% = 4 of 10, exact

    builder = _builder(plan, facts, tables)
    content = _Content(SlideType.KPI, plan.slides[5].title, kpi_fact_ids=plan.slides[5].kpi_fact_ids)
    attempt = builder._evaluate(builder.layouts["matn:kpi:progress"], content, (0, 2, 0), None)
    bars = {d.name: d.rect for d in attempt.decor if d.name.startswith("bar_")}
    assert abs(bars["bar_fill_1"].w - 0.18 * bars["bar_track_1"].w) <= 1
    assert abs(bars["bar_fill_2"].w - 0.72 * bars["bar_track_2"].w) <= 1


def _builder(plan, facts, tables, template=DEFAULT):
    builder = _DeckBuilder(Presentation(str(template)), facts, tables, load_labels(), Path(template))
    builder.title_pt = builder._deck_title_size(plan)
    builder.sections = 2
    return builder


def test_titles_use_sentence_case_and_balanced_breaks(specimen):
    plan, facts, _, prs, report = specimen
    long_slide = next(e["slide"] - 1 for e in report.scores if e["plan_index"] == 7)  # the 90-character title
    title = by_name(prs.slides[long_slide], "title")
    runs = title.text_frame.paragraphs[0].runs
    assert all(r.font._rPr.get("cap") == "none" for r in runs)  # Integral sets cap="all" on titles
    lines = title.text_frame.text.split("\v")
    assert len(lines) >= 2 and " ".join(lines) == plan.slides[7].title
    assert len(lines[-1].split()) > 1  # no orphan word
    short = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
        Slide(slide_type=SlideType.LIMITATIONS, title="Two limits", items=(item("Sample caveat.", facts[11]),),
              notes=None),))
    prs, _ = build_presentation(short, facts, (), DEFAULT)
    assert by_name(prs.slides[0], "title").text_frame.paragraphs[0].runs[0].font._rPr.get("cap") is None


def test_no_two_consecutive_content_slides_share_a_variant():
    plan, facts, tables = specimen_inputs()
    for template in gallery():
        _, report = build_presentation(plan, facts, tables, template.path, strict=True)
        entries = report.scores
        for a, b in zip(entries, entries[1:]):
            content = {t.value for t in CONTENT_TYPES}
            if a["slide_type"] in content and b["slide_type"] in content and a["plan_index"] != b["plan_index"]:
                assert a["chosen"]["layout"] != b["chosen"]["layout"], (template.id, a["slide"])


def test_the_tracker_marks_the_current_section(specimen):
    plan, _, _, prs, report = specimen
    for entry in report.scores:
        slide = prs.slides[entry["slide"] - 1]
        dots = [s for s in slide.shapes if s.name.startswith("tracker_")]
        if entry["slide_type"] not in {t.value for t in CONTENT_TYPES} or entry["plan_index"] < 2:
            assert not dots, entry["slide"]
            continue
        section = 1 if entry["plan_index"] < 14 else 2  # plan slide 14 is the second divider
        assert [fill_slot(d) == "accent1" for d in dots] == [k == section for k in (1, 2)], entry["slide"]


def test_split_slides_never_show_an_empty_card():
    """2b regression: a split implications slide kept the 4-card layout and showed empty cards."""
    long = "a sample action described in one sentence of moderate length"  # 2x2 cannot hold four in Ion
    facts = tuple(sample_fact(f"Sample: action {i}.", f"Sample text: action {i}.", i, kind=FactKind.RECOMMENDATION)
                  for i in range(1, 5))
    plan = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
        Slide(slide_type=SlideType.IMPLICATIONS, title="Sample actions",
              items=tuple(item(f"Action {i}: {long}.", facts[i - 1]) for i in range(1, 5)), notes=None),))
    narrow = GALLERY_DIR / "ion-boardroom" / "template.pptx"
    prs, report = build_presentation(plan, facts, (), narrow, strict=True)
    assert len(prs.slides) == 2 and set(report.layouts) == {"matn:implications:two-split"}
    for slide in prs.slides:
        cards = [s for s in slide.shapes if s.name.startswith(("card_", "hero_card"))]
        texts = [s for s in slide.placeholders if s.name.startswith(("item_", "quote", "kpi_"))]
        assert len(cards) == 2
        for card in cards:
            box = Rect(card.left, card.top, card.width, card.height)
            assert any(box.contains(Rect(t.left, t.top, t.width, t.height), 12700) and t.text_frame.text.strip()
                       for t in texts), (slide.slide_layout.name, card.name)
    # the 1- and 3-item implications variants are the other split fallbacks
    for count, layout in ((1, "matn:implications:single-hero"), (3, "matn:implications:cards-3")):
        short = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
            Slide(slide_type=SlideType.IMPLICATIONS, title="Sample actions",
                  items=tuple(item(f"Action {i}: sample.", facts[i - 1]) for i in range(1, count + 1)), notes=None),))
        assert build_presentation(short, facts, (), DEFAULT, strict=True)[1].layouts == (layout,)


def test_overflow_rejects_candidates_then_splits_with_part_titles():
    numbers = tuple(sample_fact(f"Sample: n{i}.", f"Sample text: {i},000 sample units were counted.", i % 20,
                                number=(f"Sample count {i}", f"{i},000", "units")) for i in range(1, 31))
    plan = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=tuple(
        Slide(slide_type=SlideType.LIMITATIONS, title="Sample caveat", items=(item("Sample.", n),), notes=None)
        for n in numbers) + (Slide(slide_type=SlideType.EVIDENCE_APPENDIX, title="Sample receipts", notes=None),))
    prs, report = build_presentation(plan, numbers, (), DEFAULT, strict=True)
    parts = [i for i, layout in enumerate(report.layouts) if layout == "matn:evidence_appendix:default"]
    assert len(parts) >= 2
    titles = [by_name(prs.slides[i], "title").text_frame.text.replace("\v", " ") for i in parts]
    assert titles == [f"Sample receipts ({k}/{len(parts)})" for k in range(1, len(parts) + 1)]
    lines = [p.text for i in parts for p in by_name(prs.slides[i], "body").text_frame.paragraphs]
    assert len(lines) == len(numbers)  # every number once, none cut
    assert all(f"“{n.evidence[0].quote}”" in line for n, line in zip(numbers, lines))


def test_the_same_plan_keeps_identical_facts_values_and_sources_in_every_template():
    plan, facts, tables = specimen_inputs()

    def content(template: Path):
        prs, _ = build_presentation(plan, facts, tables, template, strict=True)
        notes, values, pages, appendix = set(), [], set(), []
        for slide in prs.slides:
            notes |= {line for line in slide.notes_slide.notes_text_frame.text.splitlines() if "“" in line}
            for shape in slide.placeholders:
                if shape.name.startswith(("kpi_value_", "kpi_label_")) and slide.slide_layout.name.startswith("matn:kpi"):
                    values.append(shape.text_frame.text)
                elif shape.name == "source":
                    pages |= set(re.findall(r"\d+(?: \(table\))?", shape.text_frame.text))
                elif shape.name == "chart":
                    values.append(str(sorted(list(s.values) for s in shape.chart.plots[0].series)))
            if slide.slide_layout.name == "matn:evidence_appendix:default":
                appendix += [p.text for p in by_name(slide, "body").text_frame.paragraphs]
        return notes, sorted(values), pages, appendix

    expected = content(DEFAULT)
    assert all(expected)
    for template in gallery():
        assert content(template.path) == expected, template.id


def test_rejected_facts_unknown_tables_and_invalid_templates_stop_the_render(tmp_path):
    rejected = sample_fact("Sample: rejected.", "Sample text: rejected.", 1, status=FactStatus.REJECTED)
    plan = DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(
        Slide(slide_type=SlideType.LIMITATIONS, title="Sample", items=(item("Sample.", rejected),), notes=None),))
    with pytest.raises(DeckRenderError, match="REJECTED_FACT_CITED"):
        build_presentation(plan, (rejected,), (), DEFAULT)
    broken = sample_table([["Region", "Value"], ["Region A", "about 12"]], 9)
    fact = sample_fact("Sample: a.", "Sample text: a.", 1)
    chart = Slide(slide_type=SlideType.CHART, title="Sample", items=(item("Sample.", fact),), notes=None,
                  chart=ChartSpec(table_id=broken.table_id, x_column=0, y_columns=(1,), chart_type=ChartType.BAR))
    with pytest.raises(DeckRenderError, match="UNKNOWN_TABLE_ID"):
        build_presentation(DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(chart,)), (fact,), (), DEFAULT)
    with pytest.raises(DeckRenderError, match="CHART_CELL_NOT_NUMERIC"):
        build_presentation(DeckPlan(doc_id=SPECIMEN_DOC_ID, slides=(chart,)), (fact,), (broken,), DEFAULT)
    with pytest.raises(TemplateInvalidError, match="TEMPLATE_INVALID"):
        build_presentation(plan, (fact,), (), tmp_path / "missing.pptx")


def _slides_by_layout(prs, prefix):
    return [slide for slide in prs.slides if slide.slide_layout.name.startswith(prefix)]


def test_cards_hug_their_content_and_share_a_row_height():
    plan, facts, tables = specimen_inputs()
    builder = _builder(plan, facts, tables)
    s = plan.slides[7]  # three items
    content = _Content(s.slide_type, s.title, s.items, section=1)
    for anchor in (0, 1):
        attempt = builder._evaluate(builder.layouts["matn:finding:cards-3"], content, (1, 1, anchor), None)
        cards = [d.rect for d in attempt.decor if d.name.startswith("card_")]
        assert len(cards) == 3 and len({round(c.h) for c in cards}) == 1 and len({round(c.y) for c in cards}) == 1
        texts = [attempt.rects[n] for n in attempt.fills if n.startswith("item_")]
        tallest = max(e.rect.bottom for e in attempt.elements if e.name.startswith("item_")) -             min(e.rect.y for e in attempt.elements if e.name.startswith("item_"))
        pad = 14 * 12700
        assert abs(cards[0].h - (tallest + 2 * pad)) < 12700 * 3  # content + padding, not the slot height
        assert cards[0].h < 0.8 * builder.area.h
        assert all(any(c.contains(t, 12700) for c in cards) for t in texts)
    top = builder._evaluate(builder.layouts["matn:finding:cards-3"], content, (1, 1, 0), None)
    centered = builder._evaluate(builder.layouts["matn:finding:cards-3"], content, (1, 1, 1), None)
    card_top = next(d.rect for d in top.decor if d.name == "card_1")
    card_mid = next(d.rect for d in centered.decor if d.name == "card_1")
    assert abs(card_top.y - builder.area.y) < 12700 and card_mid.y > card_top.y  # the group anchors as one


def test_one_value_kpi_slides_render_as_a_hero_number(specimen):
    plan, facts, _, prs, report = specimen
    entry = next(e for e in report.scores if e["plan_index"] == 11)
    assert entry["chosen"]["layout"] == "matn:kpi:hero-number"
    slide = prs.slides[entry["slide"] - 1]
    runs = by_name(slide, "kpi_value_1").text_frame.paragraphs[0].runs
    assert runs[0].text == "1,240" and runs[0].font.size.pt >= 96


def test_title_accents_color_only_real_numbers_at_most_two(specimen):
    plan, facts, _, prs, report = specimen
    seen = 0
    for entry in report.scores:
        slide = prs.slides[entry["slide"] - 1]
        title = by_name(slide, "title")
        plain = title.text_frame.text.replace("", " ")
        assert plain == plan.slides[entry["plan_index"]].title or entry["part"] is not None
        runs = [r for r in title.text_frame.paragraphs[0].runs]
        accent = [r.text for r in runs if r.font._rPr.find(f"{_A}solidFill") is not None]
        assert len(accent) <= 2, plain
        for text in accent:
            assert re.fullmatch(r"[+-]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?%?", text), text
        if entry["slide_type"] not in {t.value for t in CONTENT_TYPES}:
            assert not accent
        seen += bool(accent)
    assert seen >= 3
    summary = by_name(prs.slides[1], "title")
    accents = [r.text for r in summary.text_frame.paragraphs[0].runs if r.font._rPr.find(f"{_A}solidFill") is not None]
    assert accents == ["18%", "2025"]  # "18%" equals the cited growth fact; 2025 is the next number


def test_no_content_slide_uses_more_than_three_content_sizes():
    plan, facts, tables = specimen_inputs()
    for template in gallery():
        prs, report = build_presentation(plan, facts, tables, template.path, strict=True)
        for entry in report.scores:
            if entry["slide_type"] not in {t.value for t in CONTENT_TYPES}:
                continue
            slide = prs.slides[entry["slide"] - 1]
            sizes = set()
            for shape in slide.placeholders:
                if shape.name == "source":
                    continue  # chrome: a fixed meta size
                if shape.has_chart:
                    sizes.add(shape.chart.font.size.pt)
                    continue
                sizes |= {r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size}
            assert len(sizes) <= 3, (template.id, entry["slide"], sorted(sizes))
